"""Extractor de documentos de texto (0 tokens): Word, PDF, texto plano, CSV/TSV, HTML y RTF.

Devuelve el texto plano y los párrafos. No interpreta: solo extrae. La interpretación semántica
la hace un subagente después, y solo sobre el trozo que vaya a razonar.

Las tablas de un .docx se aplanan a `a | b | c` porque en la práctica muchas actas y minutas de
reunión llegan como tablas, y perder su contenido sería perder la reunión entera.
"""
from __future__ import annotations

import os
import re
from html.parser import HTMLParser
from typing import List, Tuple


def extract_document(path: str) -> dict:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".docx", ".docm"):
        parrafos = _docx(path)
    elif ext in (".csv", ".tsv"):
        parrafos = _csv(path)
    elif ext in (".html", ".htm"):
        parrafos = _html(path)
    elif ext == ".rtf":
        parrafos = _rtf(path)
    elif ext == ".pdf":
        parrafos = _pdf(path)
    else:
        # .txt, .md, .markdown, .vtt, .srt y cualquier texto sin tratamiento especial.
        parrafos = _texto_plano(path)
    texto = "\n".join(parrafos)
    return {
        "archivo": path,
        "n_parrafos": len(parrafos),
        "n_caracteres": len(texto),
        "parrafos": parrafos,
        "texto": texto,
    }


def _docx(path: str) -> List[str]:
    import docx  # python-docx
    doc = docx.Document(path)
    parrafos = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    # Tablas del docx (a veces las transcripciones/actas vienen en tablas).
    for tabla in doc.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text and c.text.strip()]
            if celdas:
                parrafos.append(" | ".join(celdas))
    return parrafos


def _texto_plano(path: str) -> List[str]:
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        return [ln.strip() for ln in fh if ln.strip()]


def _csv(path: str) -> List[str]:
    """CSV → filas como 'a | b | c' (misma convención que las tablas de .docx)."""
    import csv
    parrafos: List[str] = []
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as fh:
        muestra = fh.read(2048)
        fh.seek(0)
        try:
            dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
        except csv.Error:
            dialecto = csv.excel
        for fila in csv.reader(fh, dialecto):
            celdas = [c.strip() for c in fila if c and c.strip()]
            if celdas:
                parrafos.append(" | ".join(celdas))
    return parrafos


def _pdf(path: str) -> List[str]:
    try:
        from pypdf import PdfReader  # dependencia opcional
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # fallback
        except ImportError:
            return [f"[PDF no extraído: instala 'pypdf' para leer {os.path.basename(path)}]"]
    reader = PdfReader(path)
    parrafos: List[str] = []
    for page in reader.pages:
        texto = page.extract_text() or ""
        parrafos.extend([ln.strip() for ln in texto.splitlines() if ln.strip()])
    return parrafos


class _TextoDeHtml(HTMLParser):
    """Extrae el texto visible de un HTML descartando script, style y marcado.

    Se usa `html.parser` de la stdlib en vez de BeautifulSoup: un export de Confluence o
    SharePoint no justifica una dependencia, y aquí solo se necesita el texto.
    """

    _INVISIBLES = {"script", "style", "head", "meta", "link", "noscript"}
    _BLOQUES = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.trozos: List[str] = []
        self._omitir = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._INVISIBLES:
            self._omitir += 1
        elif tag in self._BLOQUES:
            self.trozos.append("\n")
        elif tag == "td":
            self.trozos.append(" | ")

    def handle_endtag(self, tag):
        if tag in self._INVISIBLES and self._omitir:
            self._omitir -= 1
        elif tag in self._BLOQUES:
            self.trozos.append("\n")

    def handle_data(self, data):
        if not self._omitir and data.strip():
            self.trozos.append(data.strip())


def _html(path: str) -> List[str]:
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        parser = _TextoDeHtml()
        parser.feed(fh.read())
    crudo = " ".join(parser.trozos)
    lineas = [re.sub(r"\s{2,}", " ", ln).strip(" |") for ln in crudo.split("\n")]
    return [ln for ln in lineas if ln]


# RTF: grupos de control {\...}, palabras de control \xx y escapes \'hh de la codificación.
_RTF_CONTROL = re.compile(r"\\[a-z]+-?\d*\s?|[{}]|\\'[0-9a-f]{2}", re.I)


def _rtf(path: str) -> List[str]:
    """RTF → texto. Descarta el marcado sin librería: para un acta basta y sobra."""
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        crudo = fh.read()
    texto = _RTF_CONTROL.sub(" ", crudo)
    return [ln.strip() for ln in re.split(r"\n|\\par", texto) if ln.strip()]


def extraer(ruta: str) -> Tuple[dict, str, dict]:
    """Adaptador del registro."""
    doc = extract_document(ruta)
    registro = {"n_parrafos": doc["n_parrafos"], "vista_texto": doc["texto"]}
    resumen = f"Documento: {doc['n_parrafos']} párrafos, {doc['n_caracteres']} chars"
    return registro, resumen, {}
