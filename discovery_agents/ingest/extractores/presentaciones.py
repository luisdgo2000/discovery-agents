"""Extractor de PowerPoint determinista (0 tokens).

Extrae texto de shapes, tablas y notas del orador, **y también el de las imágenes incrustadas**.
Esto último importa más de lo que parece: una presentación de arquitectura suele ser una imagen
por diapositiva, de modo que un extractor que solo lee cajas de texto devuelve un archivo casi
vacío de un insumo que en realidad es de los más densos del Discovery.

Degrada con gracia: sin `python-pptx` devuelve un aviso en lugar de fallar, y sin OCR las
imágenes quedan marcadas para el modelo de visión en vez de perderse.
"""
from __future__ import annotations

import os
from typing import List, Tuple

from .embebidos import ocr_de_incrustadas


def extract_pptx(path: str) -> dict:
    try:
        from pptx import Presentation  # python-pptx, dependencia opcional
    except ImportError:
        return {
            "archivo": path, "n_diapositivas": 0, "diapositivas": [],
            "texto": f"[PPTX no extraído: instala 'python-pptx' para leer {os.path.basename(path)}]",
            "imagenes_embebidas": 0,
        }

    prs = Presentation(path)
    diapositivas: List[dict] = []
    imagenes = 0
    for n, slide in enumerate(prs.slides, start=1):
        textos: List[str] = []
        for shape in slide.shapes:
            if shape.shape_type == 13:  # PICTURE
                imagenes += 1
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                textos.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for fila in shape.table.rows:
                    celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
                    if celdas:
                        textos.append(" | ".join(celdas))
        notas = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notas = slide.notes_slide.notes_text_frame.text.strip()
        diapositivas.append({"n": n, "texto": "\n".join(textos), "notas": notas})

    texto = "\n\n".join(
        f"[Diapositiva {d['n']}]\n{d['texto']}" + (f"\n[Notas] {d['notas']}" if d["notas"] else "")
        for d in diapositivas if d["texto"] or d["notas"]
    )

    # El contenido de las imágenes incrustadas se anexa al final, etiquetado, para que se pueda
    # citar igual que el resto pero sin confundirse con el texto que sí estaba en las cajas.
    incrustadas = ocr_de_incrustadas(path)
    if incrustadas["texto"]:
        texto = (texto + "\n\n== Texto de imágenes incrustadas ==\n\n"
                 + str(incrustadas["texto"])).strip()

    return {
        "archivo": path,
        "n_diapositivas": len(diapositivas),
        "diapositivas": diapositivas,
        "texto": texto,
        "imagenes_embebidas": max(imagenes, int(incrustadas["n_imagenes"])),
        "imagenes_leidas_por_ocr": int(incrustadas["n_con_texto"]),
        "imagenes_para_vision": list(incrustadas["requieren_vision"]),
    }


def extraer(ruta: str) -> Tuple[dict, str, dict]:
    """Adaptador del registro."""
    data = extract_pptx(ruta)
    para_vision = data.get("imagenes_para_vision", [])
    registro = {
        "n_diapositivas": data["n_diapositivas"],
        "diapositivas": data["diapositivas"],
        "imagenes_embebidas": data["imagenes_embebidas"],
        "imagenes_leidas_por_ocr": data.get("imagenes_leidas_por_ocr", 0),
        "imagenes_para_vision": para_vision,
        "vista_texto": data["texto"],
    }
    resumen = (f"PowerPoint: {data['n_diapositivas']} diapositivas, "
               f"{data['imagenes_embebidas']} imágenes embebidas")
    if data.get("imagenes_leidas_por_ocr"):
        resumen += f" ({data['imagenes_leidas_por_ocr']} leídas por OCR)"
    if para_vision:
        resumen += f", {len(para_vision)} requieren visión LLM"
    return registro, resumen, {
        "es_imagen": data["imagenes_embebidas"], "requiere_vision": len(para_vision),
    }
