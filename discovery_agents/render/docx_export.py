"""Exportación del Discovery a Word (.docx) — v4: documento fusionado, con portada e índice.

Es el entregable principal del sistema y hasta ahora era el peor presentado: sin portada, sin
índice, sin encabezado ni numeración, con la tipografía por defecto y con el contenido dicho dos
veces (la narrativa arriba, el mismo contenido otra vez en el anexo estructurado).

Esta versión arma **el mismo documento fusionado que define `render/ensamblador.py`** —una sección,
un cuerpo— y le da la forma que corresponde a un documento que se lleva a un comité: portada con el
veredicto, índice navegable, estilos propios, tablas con filete y cabecera repetida, y anexos donde
viven las matrices completas.

Dos decisiones que merecen explicación:

- **Los diagramas se describen en palabras, además de en código.** Mermaid no se renderiza dentro
  de un `.docx`: pegar el fuente deja al lector con veinte líneas de `flowchart LR` que no le dicen
  nada. Cada diagrama va acompañado de su lectura en prosa, generada del mismo dato —así no puede
  contradecirlo—, y el código queda en el anexo para quien quiera reproducirlo.
- **El markdown inline se convierte, no se borra.** `**negrita**` y `` `código` `` se traducen a
  formato real de Word: si se descartaran, el énfasis que el Redactor puso deliberadamente
  desaparecería justo en el entregable que se presenta.

Degrada con gracia: sin `python-docx` devuelve None y el resto de los entregables se generan igual.
"""
from __future__ import annotations

import os
import re
from typing import Dict, List, Optional

from ..common.citas import tiene_cita
from ..common.marcas import es_ausente
from ..knowledge.schema import (COLUMNAS_PREFERIDAS, ETIQUETA_TIPO_RIESGO, ETIQUETAS, ORDEN_PARTES,
                                PARTES, PENDIENTE, SECCIONES, TIPOS_RIESGO, render_como_tabla)
from ..knowledge.scorer import score_discovery
from . import estilos_word
from .ensamblador import _slots_de_prosa, tiene_contenido
from .mermaid import describir_diagrama

_SIN_BENEFICIO = "No cuantificado con evidencia en el Discovery."


def docx_disponible() -> bool:
    try:
        import docx  # noqa: F401
        return True
    except ImportError:
        return False


def exportar_docx(path: str, discovery: dict, caso: str = "", score: Optional[dict] = None,
                  redaccion: Optional[dict] = None, fidelidad: Optional[dict] = None,
                  diagramas: Optional[Dict[str, str]] = None,
                  secciones_omitidas: Optional[List[str]] = None) -> Optional[str]:
    """Escribe el documento Word y devuelve su ruta, o None si falta `python-docx`."""
    if not docx_disponible():
        return None

    from docx import Document

    score = score or score_discovery(discovery)
    fidelidad = fidelidad or {}
    prosa = _slots_de_prosa(redaccion)
    titulo = caso or str(discovery.get("nombre_caso_uso") or "Producto de datos")

    doc = Document()
    estilos_word.preparar_documento(doc, titulo_encabezado=f"Discovery — {titulo}")

    fecha = str(discovery.get("fecha_discovery", "")).strip()
    estilos_word.portada(
        doc, titulo=titulo, fecha="" if es_ausente(fecha) else fecha,
        veredicto=score["veredicto"], completitud=score["completitud_total"],
        cobertura=fidelidad.get("cobertura_pct", 0),
        n_casos=_n_casos(discovery), n_fuentes=len(_fuentes_validas(discovery)),
    )
    estilos_word.indice(doc)

    _resumen_ejecutivo(doc, discovery, score)

    for parte in ORDEN_PARTES:
        secciones = [s for s in SECCIONES if s["parte"] == parte]
        cuerpos = [s for s in secciones if tiene_contenido(s, discovery, prosa)]
        if not cuerpos:
            continue
        doc.add_heading(PARTES[parte], level=1)
        for seccion in cuerpos:
            doc.add_heading(str(seccion["titulo"]), level=2)
            _cuerpo_seccion(doc, seccion, discovery, score, fidelidad, prosa)

    if secciones_omitidas:
        _no_capturado(doc, secciones_omitidas)

    _anexos(doc, discovery, diagramas or {})

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    doc.save(path)
    return path


# ---------------------------------------------------------------------------------------------
# Secciones
#
# La regla de "qué sección merece imprimirse" vive en `ensamblador.tiene_contenido` y se importa
# desde allí: si el markdown y el .docx la calcularan cada uno por su lado, acabarían discrepando
# sobre qué secciones existen, y el documento dejaría de ser el mismo en los dos formatos.
# ---------------------------------------------------------------------------------------------
def _cuerpo_seccion(doc, seccion: dict, discovery: dict, score: dict,
                    fidelidad: dict, prosa: Dict[str, str]) -> None:
    texto = prosa.get(str(seccion.get("narrativa") or ""), "").strip()
    if texto:
        _md_a_docx(doc, texto)

    for clave in seccion["claves"]:
        _agregar_clave(doc, clave, discovery, score, fidelidad, ya_narrado=bool(texto))


def _agregar_clave(doc, clave: str, discovery: dict, score: dict, fidelidad: dict,
                   ya_narrado: bool) -> None:
    if clave == "_veredicto":
        _veredicto_docx(doc, score)
        return
    if clave == "_cobertura":
        _cobertura_docx(doc, fidelidad)
        return
    if clave == "_completitud_bloques":
        _bloques_docx(doc, score)
        return
    if clave == "riesgos":
        _riesgos_docx(doc, discovery.get("riesgos"))
        return
    if clave == "temas_faltantes":
        _temas_docx(doc, discovery.get("temas_faltantes"))
        return

    valor = discovery.get(clave)
    etiqueta = ETIQUETAS.get(clave, clave)

    if clave == "beneficio_economico":
        if es_ausente(valor) or not tiene_cita(valor):
            _campo_texto(doc, etiqueta, _SIN_BENEFICIO, ausente=True)
        else:
            _campo_texto(doc, etiqueta, str(valor))
        return

    if isinstance(valor, str):
        if es_ausente(valor):
            if not ya_narrado:
                _campo_texto(doc, etiqueta, PENDIENTE, ausente=True)
        else:
            _campo_texto(doc, etiqueta, valor)
        return

    if isinstance(valor, dict):
        persona = _persona_str(valor)
        if persona:
            _campo_texto(doc, etiqueta, persona)
        return

    if isinstance(valor, list) and valor:
        filas = [v for v in valor if isinstance(v, dict)
                 and any(str(x).strip() for x in v.values())]
        if filas:
            _subtitulo(doc, etiqueta)
            if render_como_tabla(clave):
                estilos_word.tabla(doc, filas, COLUMNAS_PREFERIDAS.get(clave), ETIQUETAS)
            else:
                _lista_objetos(doc, clave, filas)
        elif all(isinstance(v, str) for v in valor):
            _subtitulo(doc, etiqueta)
            for v in valor:
                if str(v).strip():
                    doc.add_paragraph(str(v).strip(), style="List Bullet")


def _riesgos_docx(doc, riesgos: object) -> None:
    """Riesgos agrupados por tipo; si no hay, se dice con precisión."""
    validos = [r for r in (riesgos or []) if isinstance(r, dict)
               and not es_ausente(r.get("descripcion"))]
    if not validos:
        parrafo = doc.add_paragraph(style="DiscoveryNota")
        parrafo.add_run(
            "No se identificaron riesgos durante el Discovery. Ausencia de evidencia de riesgo, "
            "no garantía de su inexistencia.")
        return

    por_tipo: Dict[str, List[dict]] = {}
    for r in validos:
        tipo = str(r.get("tipo", "")).strip().lower()
        por_tipo.setdefault(tipo if tipo in TIPOS_RIESGO else "", []).append(r)

    orden = [t for t in TIPOS_RIESGO if t in por_tipo] + ([""] if "" in por_tipo else [])
    for tipo in orden:
        _subtitulo(doc, ETIQUETA_TIPO_RIESGO[tipo])
        estilos_word.tabla(doc, por_tipo[tipo], COLUMNAS_PREFERIDAS.get("riesgos"), ETIQUETAS)


def _temas_docx(doc, temas: object) -> None:
    """Temas faltantes agrupados por impacto; si no hay ninguno, se dice explícitamente."""
    lista = [t for t in (temas or []) if isinstance(t, dict) and str(t.get("punto", "")).strip()]
    if not lista:
        doc.add_paragraph(
            "El Discovery respondió todas las preguntas de la rúbrica y no dejó supuestos sin "
            "validar: no hay temas faltantes ni deuda técnica registrada.", style="DiscoveryNota")
        return

    grupos = [("Bloqueante", "Bloqueantes"), ("Recomendado", "Recomendados"),
              ("Opcional", "Opcionales")]
    clasificados = {impacto for impacto, _ in grupos}
    for impacto, titulo in grupos:
        grupo = [t for t in lista if str(t.get("impacto", "")) == impacto]
        if grupo:
            _subtitulo(doc, titulo)
            _bullets_temas(doc, grupo)
    resto = [t for t in lista if str(t.get("impacto", "")) not in clasificados]
    if resto:
        _subtitulo(doc, "Otros temas abiertos")
        _bullets_temas(doc, resto)


def _bullets_temas(doc, temas: List[dict]) -> None:
    for t in temas:
        parrafo = doc.add_paragraph(style="List Bullet")
        parrafo.add_run(str(t["punto"])).bold = True
        extras = [str(t.get(k, "")).strip() for k in ("detalle", "accion_sugerida")]
        extras = [x for x in extras if x and not es_ausente(x)]
        if extras:
            parrafo.add_run(" — " + " · ".join(extras))
        origen = str(t.get("origen", "")).strip()
        if origen:
            parrafo.add_run(f" (detectado por: {origen})").italic = True


def _bloques_docx(doc, score: dict) -> None:
    filas = [{
        "bloque": b["titulo"],
        "completitud": f"{b['completitud']}%",
        "peso": f"{b['peso']}%",
        "critico": "Sí" if b["critico"] else "No",
        "pendientes": str(len(b["pendientes"])),
    } for b in score.get("bloques", [])]
    if filas:
        estilos_word.tabla(doc, filas, ["bloque", "completitud", "peso", "critico", "pendientes"],
                           {"critico": "Crítico", "pendientes": "Preguntas pendientes"})


def _veredicto_docx(doc, score: dict) -> None:
    from docx.shared import Pt, RGBColor

    parrafo = doc.add_paragraph()
    run = parrafo.add_run(
        f"Veredicto: {score['veredicto']} · Completitud {score['completitud_total']}% · "
        f"{score['n_brechas']} brecha(s)")
    run.font.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(
        estilos_word.COLOR_VEREDICTO.get(score["veredicto"], estilos_word.PALETA["condicional"]))

    condiciones = score.get("condiciones_para_avanzar") or []
    if condiciones:
        _subtitulo(doc, "Condiciones para avanzar")
        for c in condiciones:
            doc.add_paragraph(str(c), style="List Bullet")
    else:
        doc.add_paragraph("Sin condiciones bloqueantes pendientes.", style="DiscoveryNota")

    brechas = score.get("brechas") or []
    if brechas:
        _subtitulo(doc, "Preguntas pendientes de la rúbrica")
        for b in brechas:
            marca = ("[BLOQUEANTE] " if b["critico"] and b["obligatoria"]
                     else "[Recomendada] " if b["obligatoria"] else "[Opcional] ")
            doc.add_paragraph(f"{marca}[{b['bloque']}] {b['pregunta']}", style="List Bullet")


def _cobertura_docx(doc, fidelidad: dict) -> None:
    if not fidelidad:
        return
    _campo_texto(doc, "Cobertura de citas", f"{fidelidad.get('cobertura_pct', 0)}% "
                 f"({fidelidad.get('campos_con_cita', 0)} de "
                 f"{fidelidad.get('campos_con_contenido', 0)} campos con evidencia)")
    for etiqueta, valores in (
        ("Campos sin cita", fidelidad.get("campos_sin_cita") or []),
        ("Fuentes sin evidencia", fidelidad.get("fuentes_sin_evidencia") or []),
    ):
        if valores:
            _campo_texto(doc, etiqueta, ", ".join(str(v) for v in valores))
    invalidas = fidelidad.get("citas_invalidas") or []
    if invalidas:
        _campo_texto(doc, "Citas que no corresponden a ninguna evidencia",
                     ", ".join(f"{c.get('campo')}: {c.get('cita')}" for c in invalidas))
    sin_cita = fidelidad.get("parrafos_sin_cita") or []
    if sin_cita:
        _campo_texto(doc, "Párrafos de la narrativa sin cita", str(len(sin_cita)))
    redundantes = fidelidad.get("parrafos_redundantes") or []
    if redundantes:
        _campo_texto(doc, "Párrafos que repiten contenido de otra sección", str(len(redundantes)))


def _resumen_ejecutivo(doc, discovery: dict, score: dict) -> None:
    texto = str(discovery.get("resumen_ejecutivo") or "").strip()
    mensajes = [m for m in (discovery.get("mensajes_clave") or []) if str(m).strip()]
    if es_ausente(texto) and not mensajes:
        return
    doc.add_heading("Resumen ejecutivo", level=1)
    if not es_ausente(texto):
        _md_a_docx(doc, texto)
    if mensajes:
        _subtitulo(doc, "Mensajes clave")
        for m in mensajes:
            doc.add_paragraph(str(m), style="List Bullet")
    estilos_word.salto_de_pagina(doc)


def _no_capturado(doc, omitidas: List[str]) -> None:
    doc.add_heading("Información no capturada en este Discovery", level=1)
    doc.add_paragraph(
        "Las siguientes secciones no se incluyen porque no se obtuvo información para ninguno de "
        "sus campos. Cada una está registrada como brecha en «Temas faltantes y deuda técnica».",
        style="DiscoveryNota")
    for titulo in omitidas:
        doc.add_paragraph(titulo, style="List Bullet")


# ---------------------------------------------------------------------------------------------
# Anexos
# ---------------------------------------------------------------------------------------------
def _anexos(doc, discovery: dict, diagramas: Dict[str, str]) -> None:
    """Las matrices completas y los diagramas: material de consulta, no de lectura."""
    estilos_word.salto_de_pagina(doc)
    doc.add_heading("Anexos", level=1)

    doc.add_heading("Anexo A — Trazabilidad de la evidencia", level=2)
    doc.add_paragraph(
        "Qué insumo respalda cada campo del Discovery. Los identificadores [E00x] corresponden al "
        "índice de evidencia del caso.", style="DiscoveryNota")
    from .tablas import filas_trazabilidad
    filas = filas_trazabilidad(discovery)
    if filas:
        estilos_word.tabla(doc, filas, ["campo", "evidencias"],
                           {"campo": "Campo", "evidencias": "Evidencias"})
    else:
        doc.add_paragraph("Sin citas de evidencia registradas.", style="DiscoveryNota")

    if diagramas:
        doc.add_heading("Anexo B — Diagramas", level=2)
        for nombre, contenido in diagramas.items():
            titulo = nombre.replace(".md", "").replace("_", " ").capitalize()
            doc.add_heading(titulo, level=3)
            lectura = describir_diagrama(nombre, discovery)
            if lectura:
                doc.add_paragraph(lectura)
            _bloque_codigo(doc, _solo_mermaid(contenido))


# ---------------------------------------------------------------------------------------------
# Primitivas de render
# ---------------------------------------------------------------------------------------------
def _subtitulo(doc, texto: str) -> None:
    doc.add_heading(texto, level=3)


def _campo_texto(doc, etiqueta: str, valor: str, ausente: bool = False) -> None:
    parrafo = doc.add_paragraph()
    marca = parrafo.add_run(f"{etiqueta}: ")
    marca.bold = True
    if ausente:
        run = parrafo.add_run(valor)
        run.italic = True
    else:
        _runs_con_formato(parrafo, valor)


def _lista_objetos(doc, clave: str, filas: List[dict]) -> None:
    """Lista de objetos como prosa: el primer campo encabeza y el resto lo matiza."""
    columnas = COLUMNAS_PREFERIDAS.get(clave) or list(filas[0].keys())
    cabeza = columnas[0]
    for fila in filas:
        parrafo = doc.add_paragraph(style="List Bullet")
        titular = str(fila.get(cabeza, "")).strip()
        if titular:
            run = parrafo.add_run(titular)
            run.bold = True
        detalles = []
        for columna in columnas[1:]:
            valor = fila.get(columna)
            if isinstance(valor, (list, tuple)):
                valor = ", ".join(str(v) for v in valor)
            valor = str(valor or "").strip()
            if valor and not es_ausente(valor) and columna != "evidencia":
                detalles.append(f"{ETIQUETAS.get(columna, columna.capitalize())}: {valor}")
        if detalles:
            parrafo.add_run(" — " + " · ".join(detalles))
        evidencia = str(fila.get("evidencia", "")).strip()
        if evidencia:
            cita = parrafo.add_run(f"  {evidencia}")
            cita.italic = True


def _persona_str(p: dict) -> str:
    nombre = str(p.get("nombre", "")).strip()
    if es_ausente(nombre):
        return ""
    rol = str(p.get("rol", "")).strip()
    contacto = str(p.get("contacto", "")).strip()
    salida = nombre
    if rol and not es_ausente(rol):
        salida += f" ({rol})"
    if contacto and not es_ausente(contacto):
        salida += f" — {contacto}"
    return salida


def _fuentes_validas(discovery: dict) -> List[dict]:
    return [f for f in discovery.get("fuentes", [])
            if isinstance(f, dict) and str(f.get("tabla", "")).strip()]


def _n_casos(discovery: dict) -> int:
    casos = [c for c in discovery.get("casos_uso", [])
             if isinstance(c, dict) and str(c.get("nombre", "")).strip()]
    return len(casos) or 1


# ---------------------------------------------------------------------------------------------
# Conversor Markdown → docx (subset que usa la prosa del Redactor)
# ---------------------------------------------------------------------------------------------
_INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)")


def _runs_con_formato(parrafo, texto: str) -> None:
    """Traduce el markdown inline a formato real de Word en vez de borrarlo.

    El conversor anterior eliminaba `**negrita**` y `` `código` ``, de modo que el énfasis que el
    Redactor había puesto deliberadamente desaparecía justo en el entregable que se presenta.
    """
    from docx.shared import Pt

    for trozo in _INLINE.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**") and len(trozo) > 4:
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("`") and trozo.endswith("`") and len(trozo) > 2:
            run = parrafo.add_run(trozo[1:-1])
            run.font.name = estilos_word.FUENTE_DATO
            run.font.size = Pt(9.5)
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            parrafo.add_run(trozo[1:-1]).italic = True
        else:
            parrafo.add_run(trozo)


def _md_a_docx(doc, md: str) -> None:
    lineas = md.splitlines()
    i = 0
    en_codigo = False
    buffer_codigo: List[str] = []
    while i < len(lineas):
        linea = lineas[i]
        s = linea.strip()
        if s.startswith("```"):
            if en_codigo:
                _bloque_codigo(doc, "\n".join(buffer_codigo))
                buffer_codigo = []
            en_codigo = not en_codigo
            i += 1
            continue
        if en_codigo:
            buffer_codigo.append(linea)
            i += 1
            continue
        if not s:
            i += 1
            continue
        if s.startswith("#"):
            nivel = min(len(s) - len(s.lstrip("#")), 3)
            doc.add_heading(s.lstrip("#").strip(), level=max(3, nivel + 2))
            i += 1
            continue
        if s.startswith(("- ", "* ")):
            parrafo = doc.add_paragraph(style="List Bullet")
            _runs_con_formato(parrafo, s[2:].strip())
            i += 1
            continue
        if s.startswith(">"):
            parrafo = doc.add_paragraph(style="DiscoveryNota")
            parrafo.add_run(s.lstrip(">").strip())
            i += 1
            continue
        if s.startswith("|"):
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                bloque.append(lineas[i].strip())
                i += 1
            _tabla_md_a_docx(doc, bloque)
            continue
        parrafo = doc.add_paragraph()
        _runs_con_formato(parrafo, s)
        i += 1


def _tabla_md_a_docx(doc, bloque: List[str]) -> None:
    filas = []
    for fila in bloque:
        celdas = [c.strip() for c in fila.strip().strip("|").split("|")]
        if all(set(c) <= {"-", ":", " "} for c in celdas):  # fila separadora ---
            continue
        filas.append(celdas)
    if not filas:
        return
    columnas = filas[0]
    registros = [dict(zip(columnas, fila)) for fila in filas[1:]]
    if registros:
        estilos_word.tabla(doc, registros, columnas)


def _bloque_codigo(doc, texto: str) -> None:
    if not texto.strip():
        return
    parrafo = doc.add_paragraph(style="DiscoveryDato")
    parrafo.add_run(texto.strip())


def _solo_mermaid(md: str) -> str:
    m = re.search(r"```mermaid\n(.*?)```", md, re.DOTALL)
    return m.group(1).strip() if m else md.strip()
