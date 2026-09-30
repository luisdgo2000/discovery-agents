"""Sistema de estilos del documento Word (0 tokens, sin plantilla binaria).

El Word se generaba con `Document()` a secas: Calibri 11 por defecto, sin portada, sin índice, sin
encabezado, sin numeración de páginas y con las tablas en un estilo prestado. Era el entregable
principal del sistema y el peor presentado de todos — un documento que se lleva a un comité no
puede parecer un volcado.

Aquí los estilos se construyen **por código y no desde un .docx de plantilla**. La diferencia
importa: un binario de plantilla no se puede revisar en un diff, no se puede probar y se convierte
en un archivo que nadie se atreve a tocar. Un módulo de estilos se lee, se versiona y se prueba.

La paleta es la misma del Resumen Ejecutivo HTML, para que los dos entregables del Discovery se
reconozcan como piezas de la misma familia. Como allá, el único color con carga es el del
veredicto; todo lo demás son tintas y filetes.
"""
from __future__ import annotations

from typing import List, Optional

# Paleta compartida con el Resumen Ejecutivo HTML. Neutros con sesgo frío —elegidos, no heredados—
# y una sola variable cromática: el veredicto.
PALETA = {
    "tinta": "141A19",       # casi negro con sesgo verde-pizarra
    "grafito": "606B69",     # texto secundario
    "filete": "DDE2E0",      # líneas de tabla
    "fondo_cabecera": "F2F4F3",
    "go": "1F6B4A",
    "condicional": "9A6516",
    "nogo": "93332B",
}

COLOR_VEREDICTO = {"Go": PALETA["go"], "Condicional": PALETA["condicional"],
                   "No-Go": PALETA["nogo"]}

# Tipografías: serif para títulos (registro de dictamen), sans para el cuerpo, mono para los datos
# medidos. Se eligen familias presentes por defecto en Word en Windows y macOS: incrustar una
# tipografía de marca haría el archivo pesado y dependiente de una licencia.
FUENTE_TITULO = "Georgia"
FUENTE_CUERPO = "Calibri"
FUENTE_DATO = "Consolas"

_ESTILOS_PROPIOS = "DiscoveryCuerpo", "DiscoveryDato", "DiscoveryNota", "DiscoveryCita"


def preparar_documento(doc, titulo_encabezado: str = "") -> None:
    """Define estilos, márgenes, encabezado y pie del documento."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    for seccion in doc.sections:
        seccion.top_margin = Cm(2.4)
        seccion.bottom_margin = Cm(2.2)
        seccion.left_margin = Cm(2.6)
        seccion.right_margin = Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = FUENTE_CUERPO
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(PALETA["tinta"])
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25

    _estilo_titulo(doc, "Title", 26, PALETA["tinta"], espacio_antes=0, espacio_despues=6)
    _estilo_titulo(doc, "Heading 1", 18, PALETA["tinta"], espacio_antes=22, espacio_despues=8)
    _estilo_titulo(doc, "Heading 2", 14, PALETA["tinta"], espacio_antes=16, espacio_despues=6)
    _estilo_titulo(doc, "Heading 3", 11.5, PALETA["grafito"], espacio_antes=12, espacio_despues=4)

    _crear_estilo(doc, "DiscoveryCuerpo", FUENTE_CUERPO, 10.5, PALETA["tinta"])
    _crear_estilo(doc, "DiscoveryDato", FUENTE_DATO, 9, PALETA["tinta"])
    _crear_estilo(doc, "DiscoveryNota", FUENTE_CUERPO, 9, PALETA["grafito"], cursiva=True)
    _crear_estilo(doc, "DiscoveryCita", FUENTE_CUERPO, 9, PALETA["grafito"])

    if titulo_encabezado:
        cabecera = doc.sections[0].header.paragraphs[0]
        cabecera.text = titulo_encabezado
        cabecera.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in cabecera.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(PALETA["grafito"])
            run.font.name = FUENTE_CUERPO

    _pie_con_paginas(doc)


def _estilo_titulo(doc, nombre: str, tamano: float, color: str,
                   espacio_antes: float, espacio_despues: float) -> None:
    from docx.shared import Pt, RGBColor

    try:
        estilo = doc.styles[nombre]
    except KeyError:  # pragma: no cover — plantilla sin ese estilo predefinido
        return
    estilo.font.name = FUENTE_TITULO
    estilo.font.size = Pt(tamano)
    estilo.font.bold = nombre != "Title"
    estilo.font.color.rgb = RGBColor.from_string(color)
    estilo.paragraph_format.space_before = Pt(espacio_antes)
    estilo.paragraph_format.space_after = Pt(espacio_despues)
    # Un encabezado nunca debe quedar solo al final de una página.
    estilo.paragraph_format.keep_with_next = True


def _crear_estilo(doc, nombre: str, fuente: str, tamano: float, color: str,
                  cursiva: bool = False) -> None:
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Pt, RGBColor

    if nombre in [s.name for s in doc.styles]:
        estilo = doc.styles[nombre]
    else:
        estilo = doc.styles.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
    estilo.font.name = fuente
    estilo.font.size = Pt(tamano)
    estilo.font.italic = cursiva
    estilo.font.color.rgb = RGBColor.from_string(color)
    estilo.paragraph_format.space_after = Pt(6)


# ---------------------------------------------------------------------------------------------
# Campos de Word (página, total de páginas, índice)
# ---------------------------------------------------------------------------------------------
def _campo(parrafo, instruccion: str, texto_provisional: str = "") -> None:
    """Inserta un campo de Word (`PAGE`, `NUMPAGES`, `TOC`, …) en el párrafo.

    python-docx no tiene API para campos, así que se construye el XML a mano. Un campo, a
    diferencia de un texto, se recalcula al abrir el documento: por eso el índice sigue siendo
    correcto aunque el documento crezca, y la numeración de páginas no hay que mantenerla.
    """
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    inicio.set(qn("w:dirty"), "true")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruccion

    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")

    provisional = OxmlElement("w:t")
    provisional.text = texto_provisional

    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")

    run = parrafo.add_run()._r
    for elemento in (inicio, instr, separador, provisional, fin):
        run.append(elemento)


def _pie_con_paginas(doc) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    pie = doc.sections[0].footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pie.add_run("Página ")
    _campo(pie, "PAGE", "1")
    pie.add_run(" de ")
    _campo(pie, "NUMPAGES", "1")
    for run in pie.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(PALETA["grafito"])
        run.font.name = FUENTE_CUERPO


def indice(doc, titulo: str = "Contenido") -> None:
    """Inserta una tabla de contenido navegable que Word actualiza sola."""
    from docx.shared import Pt, RGBColor

    encabezado = doc.add_paragraph(titulo, style="Heading 1")
    encabezado.paragraph_format.space_before = Pt(0)
    parrafo = doc.add_paragraph()
    _campo(parrafo, r'TOC \o "1-3" \h \z \u',
           "Haz clic derecho aquí y elige «Actualizar campos» para generar el índice.")
    for run in parrafo.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(PALETA["grafito"])
    salto_de_pagina(doc)


def salto_de_pagina(doc) -> None:
    from docx.enum.text import WD_BREAK

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------------------------
# Portada
# ---------------------------------------------------------------------------------------------
def portada(doc, titulo: str, fecha: str, veredicto: str, completitud: object,
            cobertura: object, n_casos: int, n_fuentes: int) -> None:
    """Portada del Discovery: qué es, de cuándo, y en qué estado llega."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    espaciador = doc.add_paragraph()
    espaciador.paragraph_format.space_after = Pt(90)

    eyebrow = doc.add_paragraph()
    run = eyebrow.add_run("DISCOVERY DE PRODUCTO DE DATOS")
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.name = FUENTE_CUERPO
    run.font.color.rgb = RGBColor.from_string(PALETA["grafito"])
    eyebrow.paragraph_format.space_after = Pt(4)

    encabezado = doc.add_paragraph(titulo, style="Title")
    encabezado.paragraph_format.space_after = Pt(10)

    if fecha:
        linea = doc.add_paragraph(style="DiscoveryNota")
        linea.add_run(f"Fecha del Discovery: {fecha}")

    # El veredicto en portada: es lo que el lector busca primero y no debería tener que ir a
    # buscarlo al final.
    dictamen = doc.add_paragraph()
    dictamen.paragraph_format.space_before = Pt(18)
    marca = dictamen.add_run(f"Veredicto: {veredicto}")
    marca.font.size = Pt(15)
    marca.font.bold = True
    marca.font.name = FUENTE_TITULO
    marca.font.color.rgb = RGBColor.from_string(
        COLOR_VEREDICTO.get(veredicto, PALETA["condicional"]))

    medidas = doc.add_paragraph(style="DiscoveryCita")
    medidas.add_run(
        f"Completitud {completitud}%  ·  Cobertura de evidencia {cobertura}%  ·  "
        f"{n_casos} caso(s) de uso  ·  {n_fuentes} fuente(s) de datos")

    nota = doc.add_paragraph(style="DiscoveryNota")
    nota.add_run(
        "El veredicto lo calcula una rúbrica determinista sobre la información capturada, no un "
        "juicio editorial. Cada afirmación factual de este documento cita el insumo del que "
        "procede.")
    nota.alignment = WD_ALIGN_PARAGRAPH.LEFT

    salto_de_pagina(doc)


# ---------------------------------------------------------------------------------------------
# Tablas
# ---------------------------------------------------------------------------------------------
def tabla(doc, filas: List[dict], columnas: Optional[List[str]] = None,
          etiquetas: Optional[dict] = None):
    """Tabla con cabecera sombreada, filetes horizontales y repetición de cabecera entre páginas."""
    from docx.shared import Pt, RGBColor

    if not filas:
        return None
    if not columnas:
        columnas = []
        for fila in filas:
            for clave in fila:
                if clave not in columnas:
                    columnas.append(clave)

    tbl = doc.add_table(rows=1, cols=len(columnas))
    tbl.autofit = True
    _sin_bordes_verticales(tbl)

    cabecera = tbl.rows[0]
    _repetir_cabecera(cabecera)
    for i, columna in enumerate(columnas):
        celda = cabecera.cells[i]
        etiqueta = (etiquetas or {}).get(columna, columna.replace("_", " ").capitalize())
        celda.text = ""
        run = celda.paragraphs[0].add_run(etiqueta)
        run.font.bold = True
        run.font.size = Pt(9)
        run.font.name = FUENTE_CUERPO
        run.font.color.rgb = RGBColor.from_string(PALETA["tinta"])
        _sombrear(celda, PALETA["fondo_cabecera"])

    for fila in filas:
        celdas = tbl.add_row().cells
        for i, columna in enumerate(columnas):
            valor = fila.get(columna, "")
            if isinstance(valor, (list, tuple)):
                valor = ", ".join(str(v) for v in valor)
            texto = str(valor or "").strip() or "—"
            celdas[i].text = ""
            run = celdas[i].paragraphs[0].add_run(texto)
            run.font.size = Pt(9)
            run.font.name = FUENTE_CUERPO
    return tbl


def _repetir_cabecera(fila) -> None:
    """Marca la fila como cabecera para que Word la repita en cada página que ocupe la tabla."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    propiedades = fila._tr.get_or_add_trPr()
    marca = OxmlElement("w:tblHeader")
    marca.set(qn("w:val"), "true")
    propiedades.append(marca)


def _sombrear(celda, color_hex: str) -> None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:fill"), color_hex)
    celda._tc.get_or_add_tcPr().append(sombra)


def _sin_bordes_verticales(tbl) -> None:
    """Filetes horizontales finos y ningún borde vertical: la retícula estorba, la línea guía."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    propiedades = tbl._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado, ancho in (("top", 4), ("bottom", 4), ("insideH", 2),
                        ("left", 0), ("right", 0), ("insideV", 0)):
        elemento = OxmlElement(f"w:{lado}")
        elemento.set(qn("w:val"), "single" if ancho else "none")
        elemento.set(qn("w:sz"), str(ancho))
        elemento.set(qn("w:color"), PALETA["filete"])
        bordes.append(elemento)
    propiedades.append(bordes)
