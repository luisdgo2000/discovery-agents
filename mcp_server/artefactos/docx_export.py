"""Exportador del Discovery a un documento Word (.docx) — entregable final (0 tokens).

Reusa la estructura tipada (SECCIONES con nivel, COLUMNAS_PREFERIDAS, ETIQUETAS) para producir
**tablas nativas de Word** (no texto) y arma el documento en sus dos vistas:
  Parte I — Discovery inicial (general) · Parte II — Discovery técnico (profundo) · Cierre.
Incluye la narrativa del Redactor (Markdown→docx ligero), Temas faltantes, Cobertura de
evidencia y un anexo con el código Mermaid de los diagramas.

Degrada con gracia: si `python-docx` no está instalado, `exportar_docx` devuelve None y el resto
del pipeline sigue produciendo los .md como siempre.
"""
from __future__ import annotations

import re
from typing import Dict, List, Optional

from ..discovery.scorer import score_discovery
from ..discovery.template import (COLUMNAS_PREFERIDAS, ETIQUETAS, PARTES, PENDIENTE,
                                  SECCIONES, es_ausente)

_ORDEN_PARTES = ["general", "tecnico", "cierre"]
_CITA = re.compile(r"\[(E\d{3}[^\]]*)\]")


def docx_disponible() -> bool:
    try:
        import docx  # noqa: F401
        return True
    except Exception:
        return False


def exportar_docx(path: str, discovery: dict, caso: str = "",
                  score: Optional[dict] = None, redaccion_md: Optional[str] = None,
                  fidelidad: Optional[dict] = None,
                  diagramas: Optional[Dict[str, str]] = None) -> Optional[str]:
    """Genera el .docx en `path`. Devuelve la ruta, o None si python-docx no está disponible."""
    try:
        from docx import Document
        from docx.shared import Pt
    except Exception:
        return None

    score = score or score_discovery(discovery)
    doc = Document()

    doc.add_heading(f"Discovery — {caso or 'Producto de datos'}", level=0)
    p = doc.add_paragraph()
    r = p.add_run(f"Veredicto: {score['veredicto']}  ·  Completitud: "
                  f"{score['completitud_total']}%  ·  Brechas: {score['n_brechas']}")
    r.bold = True

    # Resumen de completitud por bloque.
    doc.add_heading("Evaluación por bloque", level=2)
    _tabla_docx(doc, [
        {"bloque": b["titulo"], "completitud": f"{b['completitud']}%",
         "crítico": "Sí" if b["critico"] else "No"}
        for b in score["bloques"]
    ], ["bloque", "completitud", "crítico"])

    if redaccion_md and redaccion_md.strip():
        doc.add_heading("Narrativa del Discovery", level=1)
        _md_a_docx(doc, redaccion_md.strip())

    hay_narrativa = bool(redaccion_md and redaccion_md.strip())
    for nivel in _ORDEN_PARTES:
        secs = [s for s in SECCIONES if s.get("nivel") == nivel]
        if not secs:
            continue
        doc.add_heading(PARTES[nivel], level=1)
        for sec in secs:
            sec_id = sec["id"]
            if hay_narrativa and sec_id == "resumen_ejecutivo":
                continue
            doc.add_heading(str(sec["titulo"]), level=2)
            if sec_id == "veredicto":
                _veredicto_docx(doc, score); continue
            if sec_id == "brechas":
                _brechas_docx(doc, score); continue
            if sec_id == "cobertura":
                _cobertura_docx(doc, fidelidad); continue
            for clave in sec["claves"]:  # type: ignore[union-attr]
                _agregar_valor(doc, clave, discovery.get(clave))

    if diagramas:
        doc.add_heading("Anexo — Diagramas (Mermaid)", level=1)
        doc.add_paragraph("Código Mermaid (renderizable en GitHub, VS Code o SharePoint moderno):")
        for nombre, contenido in diagramas.items():
            doc.add_heading(nombre.replace("_", " ").replace(".md", "").capitalize(), level=2)
            _bloque_codigo(doc, _solo_mermaid(contenido), Pt)

    doc.save(path)
    return path


# ---------------------------------------------------------------------------------------------
# Render por tipo de campo
# ---------------------------------------------------------------------------------------------
def _agregar_valor(doc, clave: str, valor: object) -> None:
    etiqueta = ETIQUETAS.get(clave, clave)
    if isinstance(valor, list):
        if valor and isinstance(valor[0], dict):
            doc.add_paragraph().add_run(f"{etiqueta}:").bold = True
            _tabla_docx(doc, valor, COLUMNAS_PREFERIDAS.get(clave))
            return
        if valor:
            doc.add_paragraph().add_run(f"{etiqueta}:").bold = True
            for v in valor:
                doc.add_paragraph(str(v), style="List Bullet")
            return
        _campo_texto(doc, etiqueta, PENDIENTE)
        return
    if isinstance(valor, dict):
        _campo_texto(doc, etiqueta, _persona_str(valor))
        return
    _campo_texto(doc, etiqueta, str(valor or ""))


def _persona_str(valor: dict) -> str:
    partes = [str(valor.get("nombre", "")).strip()]
    if str(valor.get("rol", "")).strip():
        partes.append(f"({valor['rol']})")
    if str(valor.get("contacto", "")).strip():
        partes.append(f"— {valor['contacto']}")
    texto = " ".join(p for p in partes if p)
    if texto:
        return texto
    otros = ", ".join(f"{k}: {v}" for k, v in valor.items() if str(v).strip())
    return otros or PENDIENTE


def _campo_texto(doc, etiqueta: str, texto: str) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{etiqueta}: ").bold = True
    valor = (texto or "").strip()
    run = p.add_run(PENDIENTE if es_ausente(valor) else valor)
    if es_ausente(valor):
        run.italic = True


def _tabla_docx(doc, filas: List[dict], columnas: Optional[List[str]] = None) -> None:
    filas = [f for f in filas if isinstance(f, dict) and any(str(v).strip() for v in f.values())]
    if not filas:
        p = doc.add_paragraph()
        p.add_run(PENDIENTE).italic = True
        return
    if not columnas:
        columnas = []
        for f in filas:
            for k in f:
                if k not in columnas:
                    columnas.append(k)
    tabla = doc.add_table(rows=1, cols=len(columnas))
    try:
        tabla.style = "Light Grid Accent 1"
    except Exception:
        pass
    hdr = tabla.rows[0].cells
    for i, c in enumerate(columnas):
        hdr[i].text = c.replace("_", " ").capitalize()
        for par in hdr[i].paragraphs:
            for run in par.runs:
                run.bold = True
    for f in filas:
        celdas = tabla.add_row().cells
        for i, c in enumerate(columnas):
            v = f.get(c, "")
            if isinstance(v, list):
                v = ", ".join(str(x) for x in v)
            celdas[i].text = str(v).strip() or "—"


def _veredicto_docx(doc, score: dict) -> None:
    cond = score["condiciones_para_avanzar"]
    if cond:
        doc.add_paragraph().add_run("Condiciones para avanzar:").bold = True
        for c in cond:
            doc.add_paragraph(c, style="List Bullet")
    else:
        doc.add_paragraph("Sin condiciones bloqueantes pendientes.")


def _brechas_docx(doc, score: dict) -> None:
    if not score["brechas"]:
        doc.add_paragraph("Sin brechas: el discovery responde todas las preguntas de la rúbrica.")
        return
    for br in score["brechas"]:
        sev = "BLOQUEANTE" if (br["critico"] and br["obligatoria"]) else \
              "Recomendada" if br["obligatoria"] else "Opcional"
        doc.add_paragraph(f"[{sev}] [{br['bloque']}] {br['pregunta']}", style="List Bullet")


def _cobertura_docx(doc, fidelidad: Optional[dict]) -> None:
    if not fidelidad:
        doc.add_paragraph("Sin verificación de fidelidad disponible.")
        return
    p = doc.add_paragraph()
    p.add_run(f"Cobertura de evidencia: {fidelidad.get('cobertura_pct', 0)}% "
              f"({fidelidad.get('campos_con_cita', 0)}/{fidelidad.get('campos_con_contenido', 0)} "
              f"campos con cita [E00x]).").bold = True
    if fidelidad.get("campos_sin_cita"):
        doc.add_paragraph("Campos con contenido pero sin cita de evidencia:")
        for c in fidelidad["campos_sin_cita"]:
            doc.add_paragraph(str(c), style="List Bullet")
    if fidelidad.get("fuentes_sin_evidencia"):
        doc.add_paragraph("Fuentes sin evidencia declarada: "
                          + ", ".join(fidelidad["fuentes_sin_evidencia"]))
    if fidelidad.get("citas_invalidas"):
        doc.add_paragraph("Citas a evidencias inexistentes (revisar):")
        for c in fidelidad["citas_invalidas"]:
            doc.add_paragraph(f"{c['campo']}: [{c['cita']}]", style="List Bullet")
    if fidelidad.get("parrafos_sin_cita"):
        doc.add_paragraph(f"Párrafos de la narrativa sin cita: "
                          f"{len(fidelidad['parrafos_sin_cita'])}.")


# ---------------------------------------------------------------------------------------------
# Conversor Markdown → docx (subset usado por la narrativa del Redactor)
# ---------------------------------------------------------------------------------------------
def _md_a_docx(doc, md: str) -> None:
    lineas = md.splitlines()
    i = 0
    en_codigo = False
    buffer_codigo: List[str] = []
    while i < len(lineas):
        linea = lineas[i]
        s = linea.strip()
        if s.startswith("```"):
            if en_codigo:
                _bloque_codigo(doc, "\n".join(buffer_codigo), None)
                buffer_codigo = []
            en_codigo = not en_codigo
            i += 1
            continue
        if en_codigo:
            buffer_codigo.append(linea)
            i += 1
            continue
        if not s:
            i += 1
            continue
        if s.startswith("#"):
            nivel = min(len(s) - len(s.lstrip("#")), 4)
            doc.add_heading(s.lstrip("#").strip(), level=max(2, nivel + 1))
            i += 1
            continue
        if s.startswith(("- ", "* ")):
            doc.add_paragraph(s[2:].strip(), style="List Bullet")
            i += 1
            continue
        if s.startswith(">"):
            p = doc.add_paragraph(s.lstrip(">").strip())
            for run in p.runs:
                run.italic = True
            i += 1
            continue
        if s.startswith("|"):
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                bloque.append(lineas[i].strip())
                i += 1
            _tabla_md_a_docx(doc, bloque)
            continue
        doc.add_paragraph(_limpiar_inline(s))
        i += 1


def _tabla_md_a_docx(doc, bloque: List[str]) -> None:
    filas = []
    for fila in bloque:
        celdas = [c.strip() for c in fila.strip().strip("|").split("|")]
        if all(set(c) <= {"-", ":", " "} for c in celdas):  # fila separadora ---
            continue
        filas.append(celdas)
    if not filas:
        return
    ncol = max(len(f) for f in filas)
    tabla = doc.add_table(rows=0, cols=ncol)
    try:
        tabla.style = "Light Grid Accent 1"
    except Exception:
        pass
    for idx, fila in enumerate(filas):
        celdas = tabla.add_row().cells
        for j in range(ncol):
            celdas[j].text = _limpiar_inline(fila[j]) if j < len(fila) else ""
            if idx == 0:
                for par in celdas[j].paragraphs:
                    for run in par.runs:
                        run.bold = True


def _bloque_codigo(doc, texto: str, Pt) -> None:
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.font.name = "Courier New"
    try:
        from docx.shared import Pt as _Pt
        run.font.size = _Pt(9)
    except Exception:
        pass


def _limpiar_inline(texto: str) -> str:
    """Quita marcas inline básicas de Markdown (**negrita**, `código`)."""
    texto = re.sub(r"\*\*(.+?)\*\*", r"\1", texto)
    texto = re.sub(r"`(.+?)`", r"\1", texto)
    return texto.strip()


def _solo_mermaid(md: str) -> str:
    m = re.search(r"```mermaid\n(.*?)```", md, re.DOTALL)
    return m.group(1).strip() if m else md.strip()
