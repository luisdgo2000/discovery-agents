"""Extractor de documentos determinista (0 tokens): .docx, .txt, .md y .pdf (si hay lib).

Devuelve el texto plano y los párrafos. No interpreta: solo extrae. La interpretación semántica
la hace el LLM de Copilot en una fase posterior y solo si hace falta.
"""
from __future__ import annotations

import os
from typing import List


def extract_document(path: str) -> dict:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        parrafos = _docx(path)
    elif ext in (".txt", ".md", ".vtt", ".srt"):
        parrafos = _texto_plano(path)
    elif ext == ".pdf":
        parrafos = _pdf(path)
    else:
        parrafos = []
    texto = "\n".join(parrafos)
    return {
        "archivo": path,
        "n_parrafos": len(parrafos),
        "n_caracteres": len(texto),
        "parrafos": parrafos,
        "texto": texto,
    }


def _docx(path: str) -> List[str]:
    import docx  # python-docx
    doc = docx.Document(path)
    parrafos = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    # Tablas del docx (a veces las transcripciones/actas vienen en tablas).
    for tabla in doc.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text and c.text.strip()]
            if celdas:
                parrafos.append(" | ".join(celdas))
    return parrafos


def _texto_plano(path: str) -> List[str]:
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        return [ln.strip() for ln in fh if ln.strip()]


def _pdf(path: str) -> List[str]:
    try:
        from pypdf import PdfReader  # dependencia opcional
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # fallback
        except ImportError:
            return [f"[PDF no extraído: instala 'pypdf' para leer {os.path.basename(path)}]"]
    reader = PdfReader(path)
    parrafos: List[str] = []
    for page in reader.pages:
        texto = page.extract_text() or ""
        parrafos.extend([ln.strip() for ln in texto.splitlines() if ln.strip()])
    return parrafos
