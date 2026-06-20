"""Extractor de Excel determinista (0 tokens).

Reutiliza el patrón openpyxl ya probado en el proyecto de Flujo de Dominios. Extrae celdas no
vacías y, de forma heurística, detecta tablas de FUENTES/VOLUMETRÍA que son el insumo clave del
Discovery técnico (tablas, servidor de origen, volumetría, campos).
"""
from __future__ import annotations

import re
from typing import Dict, List, Optional

import openpyxl

# Palabras clave de encabezado → campo canónico de una fila de "fuente de datos".
ENCABEZADOS_FUENTE = {
    "tabla": ["tabla", "table", "objeto", "entidad"],
    "esquema": ["esquema", "schema"],
    "base_datos": ["base de datos", "database", "bd", "db", "catalogo", "catálogo"],
    "servidor": ["servidor", "server", "origen", "instancia", "host"],
    "volumetria": ["volumetria", "volumetría", "volumen", "registros", "filas", "rows", "conteo"],
    "frecuencia": ["frecuencia", "actualizacion", "actualización", "periodicidad", "latencia"],
    "campos": ["campos", "columnas", "fields", "atributos"],
    "sensible": ["sensible", "privado", "pii", "confidencial", "regulado"],
    "descripcion": ["descripcion", "descripción", "detalle", "comentario"],
}


def _norm(s: object) -> str:
    return re.sub(r"\s+", " ", str(s).strip().lower()) if s is not None else ""


def _mapear_encabezados(fila: List[object]) -> Optional[Dict[int, str]]:
    """Si la fila parece un encabezado de tabla de fuentes, devuelve {col_idx: campo_canonico}."""
    mapa: Dict[int, str] = {}
    for idx, val in enumerate(fila):
        texto = _norm(val)
        if not texto:
            continue
        for canonico, claves in ENCABEZADOS_FUENTE.items():
            if any(k in texto for k in claves):
                mapa[idx] = canonico
                break
    # Consideramos encabezado válido si reconoce al menos "tabla" + otro campo.
    if "tabla" in mapa.values() and len(mapa) >= 2:
        return mapa
    return None


def extract_excel(path: str) -> dict:
    """Extrae contenido estructurado de un Excel.

    Returns dict con:
      - hojas: [{nombre, max_fila, max_col, celdas:[{coord, valor}]}]
      - fuentes_detectadas: [{tabla, servidor, volumetria, ...}] (heurístico)
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    hojas: List[dict] = []
    fuentes: List[dict] = []

    for ws in wb.worksheets:
        celdas: List[dict] = []
        filas_matriz: List[List[object]] = []
        for row in ws.iter_rows():
            fila_vals = [c.value for c in row]
            filas_matriz.append(fila_vals)
            for c in row:
                if c.value is not None and str(c.value).strip() != "":
                    celdas.append({"coord": c.coordinate, "valor": str(c.value).strip()})
        hojas.append({
            "nombre": ws.title,
            "max_fila": ws.max_row,
            "max_col": ws.max_column,
            "celdas": celdas,
        })
        fuentes.extend(_detectar_fuentes(ws.title, filas_matriz))

    wb.close()
    return {"archivo": path, "hojas": hojas, "fuentes_detectadas": fuentes}


def _detectar_fuentes(hoja: str, filas: List[List[object]]) -> List[dict]:
    """Recorre filas buscando un encabezado de tabla de fuentes y mapea las filas siguientes."""
    encontrados: List[dict] = []
    mapa: Optional[Dict[int, str]] = None
    for fila in filas:
        if mapa is None:
            mapa = _mapear_encabezados(fila)
            continue
        # Fila vacía corta la tabla.
        if all((v is None or str(v).strip() == "") for v in fila):
            mapa = None
            continue
        registro: Dict[str, str] = {"hoja": hoja}
        for col_idx, campo in mapa.items():
            if col_idx < len(fila) and fila[col_idx] is not None:
                valor = str(fila[col_idx]).strip()
                if valor:
                    registro[campo] = valor
        if registro.get("tabla"):
            encontrados.append(registro)
    return encontrados
