"""Extractor de PowerPoint determinista (0 tokens).

Extrae texto de shapes, tablas y notas del orador. Las imágenes embebidas se reportan para que
el flujo de imágenes decida (OCR o visión). Degrada con gracia si python-pptx no está instalado.
"""
from __future__ import annotations

import os
from typing import List


def extract_pptx(path: str) -> dict:
    try:
        from pptx import Presentation  # python-pptx, dependencia opcional
    except ImportError:
        return {
            "archivo": path, "n_diapositivas": 0, "diapositivas": [],
            "texto": f"[PPTX no extraído: instala 'python-pptx' para leer {os.path.basename(path)}]",
            "imagenes_embebidas": 0,
        }

    prs = Presentation(path)
    diapositivas: List[dict] = []
    imagenes = 0
    for n, slide in enumerate(prs.slides, start=1):
        textos: List[str] = []
        for shape in slide.shapes:
            if shape.shape_type == 13:  # PICTURE
                imagenes += 1
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                textos.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for fila in shape.table.rows:
                    celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
                    if celdas:
                        textos.append(" | ".join(celdas))
        notas = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notas = slide.notes_slide.notes_text_frame.text.strip()
        diapositivas.append({"n": n, "texto": "\n".join(textos), "notas": notas})

    texto = "\n\n".join(
        f"[Diapositiva {d['n']}]\n{d['texto']}" + (f"\n[Notas] {d['notas']}" if d["notas"] else "")
        for d in diapositivas if d["texto"] or d["notas"]
    )
    return {
        "archivo": path,
        "n_diapositivas": len(diapositivas),
        "diapositivas": diapositivas,
        "texto": texto,
        "imagenes_embebidas": imagenes,
    }
