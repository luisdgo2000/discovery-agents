"""Pruebas del núcleo determinista v2 (unittest stdlib → corre en Python 3.9+ sin pytest)."""
import json
import os
import sys
import shutil
import tempfile
import unittest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEMO = os.path.join(RAIZ, "inputs", "caso_demo")

# Workspace temporal para TODAS las pruebas (no ensuciar work/ real).
_TMP = tempfile.mkdtemp(prefix="discovery_test_")
os.environ["DISCOVERY_WORK_ROOT"] = _TMP

from discovery_agents import workspace
from discovery_agents.knowledge import consolidador as consolidator  # noqa: E402  (tras fijar el env)
from discovery_agents.render.docx_export import docx_disponible, exportar_docx  # noqa: E402
from discovery_agents.render.mermaid import generar_diagramas  # noqa: E402
from discovery_agents.render.tablas import generar_tablas, matriz_trazabilidad, render_clave  # noqa: E402
from discovery_agents.render.html_export import render_html  # noqa: E402
from discovery_agents.render.entregables import construir_entregables  # noqa: E402
from discovery_agents.ingest.extractores.documentos import extract_document  # noqa: E402
from discovery_agents.render.ensamblador import ensamblar_discovery, recolectar_temas_faltantes  # noqa: E402
from discovery_agents.knowledge.fidelidad import verificar_fidelidad  # noqa: E402
from discovery_agents.knowledge.scorer import score_discovery  # noqa: E402
from discovery_agents.knowledge.schema import schema_vacio  # noqa: E402
from discovery_agents.ingest.extractores.excel import extract_excel  # noqa: E402
from discovery_agents.ingest.manifest import build_manifest  # noqa: E402
from discovery_agents.ingest.pipeline import construir_evidencia  # noqa: E402
from discovery_agents.prompts import cargar_prompt, listar_prompts  # noqa: E402
from discovery_agents.ingest.transcripciones.cleaner import clean_transcript  # noqa: E402


def _discovery_completo():
    """Discovery tipado completo, para pruebas de Go."""
    d = schema_vacio()
    for k, v in d.items():
        if isinstance(v, str):
            d[k] = "valor de prueba [E001]"
    d["solicitante"] = {"nombre": "María Solís", "rol": "Gerente", "contacto": "msolis@x.com"}
    d["dueno_producto"] = {"nombre": "Carlos Vega", "rol": "Dueño", "contacto": "cvega@x.com"}
    d["product_owner"] = {"nombre": "Carlos Vega", "rol": "PO", "contacto": "cvega@x.com"}
    d["lider_tecnico"] = {"nombre": "Juan Pérez", "rol": "Líder técnico", "contacto": "jp@x.com"}
    d["casos_uso"] = [{"nombre": "Reportería diaria", "descripcion": "x", "evidencia": "E001"}]
    d["hallazgos_tecnicos"] = [{"hallazgo": "H", "impacto": "I", "evidencia": "E001"}]
    d["productos_relacionados"] = [{"nombre": "P", "relacion": "r", "evidencia": "E001"}]
    d["equipo_desarrollador"] = [{"nombre": "Luis Mora", "rol": "Data Engineer", "contacto": "lm@x.com"}]
    d["expertos_datos"] = [{"nombre": "Juan Pérez", "especialidad": "SQL", "contacto": "jp@x.com"}]
    d["expertos_por_pais"] = [{"nombre": "Sofía Díaz", "pais": "CRI", "especialidad": "Pagos", "contacto": "sd@x.com"}]
    d["certificadores"] = [{"nombre": "Ana Rojas", "ambito": "REG"}]
    d["fuentes"] = [
        {"tabla": "TRX", "esquema": "dbo", "servidor": "SQLPROD01",
         "columnas": ["id", "monto", "fecha"], "librerias": "pyodbc",
         "volumetria": "400M/año", "frecuencia": "Diaria", "sensible": "No", "evidencia": "E001"},
    ]
    d["entidades"] = [{"nombre": "Transaccion", "descripcion": "x", "relaciones": ["Comercio"]},
                      {"nombre": "Comercio", "descripcion": "y", "relaciones": []}]
    d["dependencias"] = [{"descripcion": "Réplica", "tipo": "DBA", "responsable": "Infra", "estado": "abierta"}]
    d["impedimentos"] = [{"descripcion": "Acceso pendiente", "responsable": "TI", "plan": "RFC", "estado": "abierto"}]
    d["accesos"] = [{"recurso": "SQLREPL02", "solicitado_a": "DBA", "estado": "en curso"}]
    d["riesgos"] = [{"descripcion": "Bloqueos", "severidad": "alta", "responsable": "Equipo"}]
    d["supuestos"] = [{"descripcion": "Volumetría vigente", "validado_por": "Carlos"}]
    d["restricciones"] = [{"descripcion": "Ventana nocturna", "origen": "Operaciones"}]
    d["hipotesis"] = [{"enunciado": "H1", "criterio_validacion": "c", "experimento": "e"}]
    d["arquetipos_datamesh"] = [{"nombre": "Plazos", "dominio": "Pagos",
                                 "valor_de_negocio": "v", "reutiliza_fuentes": ["TRX"]}]
    d["brechas"] = ""
    d["veredicto"] = ""
    return d


# =============================================================================================
# P0 — Cleaner no destructivo + formatos reales de Teams
# =============================================================================================
class TestCleanerNoDestructivo(unittest.TestCase):
    def test_preserva_palabras_con_significado(self):
        casos = [
            ("Carlos Vega: Este servidor es el que tiene la réplica.", "Este servidor"),
            ("Juan Pérez: El tipo de datos de la columna es decimal.", "tipo de datos"),
            ("María Solís: Eso no es cierto, la carga ya existe en Bronce.", "no es cierto"),
            ("María Solís: Eso no es cierto, la carga ya existe en Bronce.", "ya existe"),
            ("Carlos Vega: Bueno saber que el catálogo está listo.", "Bueno saber"),
            ("María Solís: Sí, las fuentes ya están habilitadas.", "Sí,"),
        ]
        for entrada, conservar in casos:
            out = clean_transcript(entrada)["texto_limpio"]
            self.assertIn(conservar, out, f"Se perdió '{conservar}' en: {out}")

    def test_quita_interjecciones_y_muletillas_de_inicio(self):
        out = clean_transcript("Carlos Vega: Eh, bueno, este, la idea es entender mmm el flujo, ajá.")
        texto = out["segmentos"][0]["texto"]
        self.assertEqual(texto, "la idea es entender el flujo.")

    def test_original_accesible_por_segmento(self):
        r = clean_transcript("Carlos Vega: Eh, las fuentes fuentes están listas.")
        s = r["segmentos"][0]
        self.assertEqual(s["id"], "seg-001")
        self.assertIn("fuentes fuentes", s["original"])      # original intacto
        self.assertNotIn("fuentes fuentes", s["texto"])      # limpio consolidado

    def test_dedupe_solo_adyacente(self):
        texto = ("Ana: el corte es a fin de mes.\n"
                 "Ana: el corte es a fin de mes.\n"      # glitch adyacente: se elimina
                 "Luis: ok.\n"
                 "Ana: el corte es a fin de mes.\n")     # repetición legítima lejana: se conserva
        r = clean_transcript(texto)
        todo = r["texto_limpio"].lower()
        self.assertEqual(todo.count("el corte es a fin de mes"), 2)


class TestCleanerFormatosTeams(unittest.TestCase):
    def test_vtt_teams_atribuye_hablantes(self):
        vtt = ("WEBVTT\n\n00:00:01.000 --> 00:00:04.000\n"
               "<v Carlos Vega>Las fuentes están en SQLPROD01.</v>\n\n"
               "00:00:04.500 --> 00:00:08.000\n"
               "<v Juan Pérez>Hay que replicar a otro servidor.</v>\n")
        r = clean_transcript(vtt)
        self.assertEqual(r["formato_detectado"], "vtt_teams")
        self.assertEqual(r["hablantes"], ["Carlos Vega", "Juan Pérez"])
        self.assertNotIn("<v", r["texto_limpio"])

    def test_docx_teams_no_filtra_timestamps(self):
        docx = ("Carlos Vega   0:15\nLas fuentes están en SQLPROD01.\n"
                "Juan Pérez   0:42\nPerfecto, replicamos a SQLREPL02.\n")
        r = clean_transcript(docx)
        self.assertEqual(r["formato_detectado"], "docx_teams")
        self.assertEqual(r["hablantes"], ["Carlos Vega", "Juan Pérez"])
        self.assertNotIn("0:15", r["texto_limpio"])
        self.assertNotIn("0:42", r["texto_limpio"])


# =============================================================================================
# P0 — Workspace persistente + evidencia bajo demanda
# =============================================================================================
class TestWorkspacePipeline(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        shutil.rmtree(os.path.join(_TMP, "caso_demo"), ignore_errors=True)
        cls.index = construir_evidencia(DEMO)

    def test_indice_compacto_no_token_bomb(self):
        compacto = json.dumps(self.index, ensure_ascii=False)
        self.assertLess(len(compacto), 4000, "El índice debe ser compacto (~500 tokens)")
        self.assertEqual(self.index["total_items"], 3)
        self.assertEqual(len(self.index["fuentes_detectadas"]), 3)

    def test_evidencia_persistida_con_ids(self):
        ev_dir = os.path.join(_TMP, "caso_demo", "10_evidence")
        archivos = sorted(os.listdir(ev_dir))
        self.assertIn("index.json", archivos)
        self.assertTrue(any(a.startswith("E001") for a in archivos))

    def test_get_evidence_chunk(self):
        eid = next(i["id"] for i in self.index["items"] if i["tipo"] == "transcripcion")
        ch = workspace.get_evidence_chunk("caso_demo", eid, 0)
        self.assertIn("[seg-001]", ch["contenido"])
        self.assertGreaterEqual(ch["total_chunks"], 1)
        err = workspace.get_evidence_chunk("caso_demo", "E999")
        self.assertIn("error", err)

    def test_incrementalidad_destilado(self):
        pendientes = workspace.evidencias_sin_destilar("caso_demo")
        self.assertEqual(len(pendientes), 1)  # solo la transcripción
        workspace.save_state("caso_demo", f"20_destilado/D{pendientes[0]}.json",
                             json.dumps({"decisiones": []}))
        self.assertEqual(workspace.evidencias_sin_destilar("caso_demo"), [])

    def test_save_state_lista_blanca(self):
        with self.assertRaises(ValueError):
            workspace.save_state("caso_demo", "outputs/hack.md", "x")
        with self.assertRaises(ValueError):
            workspace.save_state("caso_demo", "../fuera.json", "{}")


# =============================================================================================
# P1 — Consolidador y detector de conflictos
# =============================================================================================
class TestConsolidador(unittest.TestCase):
    def test_parse_volumen(self):
        casos = [("400,000,000 / año", 400e6), ("400M", 400e6), ("300 millones", 300e6),
                 ("18M/año", 18e6), ("120,000 registros", 120e3), ("400.000.000", 400e6),
                 ("mucho", None), ("", None)]
        for texto, esperado in casos:
            self.assertEqual(consolidator.parse_volumen(texto), esperado, f"falló: {texto!r}")

    def test_detecta_conflicto_numerico(self):
        excel = [{"tabla": "TRX_ADQUIRENCIA", "volumetria": "400,000,000 / año"}]
        reunion = [{"tabla_o_fuente": "trx adquirencia", "volumetria": "300 millones",
                    "evidencia": "E003"}]
        conflictos = consolidator.detectar_conflictos(excel, reunion)
        self.assertEqual(len(conflictos), 1)
        self.assertEqual(conflictos[0]["diferencia_pct"], 25.0)

    def test_sin_conflicto_si_coincide(self):
        excel = [{"tabla": "TRX", "volumetria": "400M"}]
        reunion = [{"tabla_o_fuente": "TRX", "volumetria": "400 millones"}]
        self.assertEqual(consolidator.detectar_conflictos(excel, reunion), [])


# =============================================================================================
# P1 — Scorer tipado con validadores
# =============================================================================================
class TestScorerTipado(unittest.TestCase):
    def test_schema_vacio_da_nogo(self):
        s = score_discovery(schema_vacio())
        self.assertEqual(s["veredicto"], "No-Go")
        self.assertEqual(s["completitud_total"], 0.0)

    def test_completo_da_go(self):
        s = score_discovery(_discovery_completo())
        self.assertEqual(s["veredicto"], "Go")
        self.assertEqual(s["completitud_total"], 100.0)

    def test_volumetria_basura_no_cuenta(self):
        d = _discovery_completo()
        d["fuentes"][0]["volumetria"] = "bastante grande"   # sin números → inválida
        s = score_discovery(d)
        claves_pendientes = [b["clave"] for b in s["brechas"]]
        self.assertIn("volumetria", claves_pendientes)
        self.assertNotEqual(s["veredicto"], "Go")

    def test_persona_sin_nombre_no_cuenta(self):
        d = _discovery_completo()
        d["dueno_producto"] = {"nombre": "", "rol": "PO", "contacto": "x"}
        s = score_discovery(d)
        self.assertIn("dueno_producto", [b["clave"] for b in s["brechas"]])

    def test_fuente_sin_servidor_no_cuenta(self):
        d = _discovery_completo()
        d["fuentes"][0]["servidor"] = ""
        s = score_discovery(d)
        self.assertIn("servidor_origen", [b["clave"] for b in s["brechas"]])


# =============================================================================================
# P1/P2 — Assembler tipado, narrativa, artefactos
# =============================================================================================
class TestAssemblerArtefactos(unittest.TestCase):
    def test_renderiza_tabla_de_fuentes(self):
        res = ensamblar_discovery(_discovery_completo(), caso_de_uso="Caso X")
        self.assertIn("| Tabla |", res["markdown"])
        self.assertIn("TRX", res["markdown"])
        self.assertIn("SQLPROD01", res["markdown"])

    def test_integra_narrativa_del_redactor(self):
        """La prosa se funde DENTRO de su sección, no en un bloque aparte."""
        res = ensamblar_discovery(
            _discovery_completo(), caso_de_uso="Caso X",
            redaccion={"secciones": {"contexto": "Narrativa del consultor [E001]."}})
        md = res["markdown"]
        self.assertIn("Narrativa del consultor", md)
        # El documento fusionado ya no tiene un bloque "Anexos — Información estructurada":
        # la prosa y los datos de la sección "Contexto" viven juntos y en ese orden.
        self.assertNotIn("Anexos — Información estructurada", md)
        pos_seccion = md.index("## Contexto y problemática actual")
        pos_prosa = md.index("Narrativa del consultor")
        pos_dato = md.index("**Problemática actual:**")
        self.assertLess(pos_seccion, pos_prosa)
        self.assertLess(pos_prosa, pos_dato)

    def test_acepta_redaccion_markdown_del_contrato_anterior(self):
        """Un 60_redaccion.md antiguo se reparte por sus encabezados y sigue funcionando."""
        res = ensamblar_discovery(
            _discovery_completo(), caso_de_uso="Caso X",
            redaccion={"markdown": "## Contexto y problema de negocio\nProsa vieja [E001]."})
        self.assertIn("Prosa vieja", res["markdown"])

    def test_ningun_dato_aparece_dos_veces(self):
        """Regla del documento fusionado: cada tabla de datos aparece una sola vez."""
        md = ensamblar_discovery(_discovery_completo(), caso_de_uso="Caso X")["markdown"]
        self.assertEqual(md.count("| Tabla | Esquema | Servidor |"), 1)
        self.assertEqual(md.count("**Objetivo del caso de uso:**"), 1)

    def test_seccion_sin_datos_se_omite(self):
        """Una sección sin prosa y sin datos no se imprime: se lista al cierre."""
        d = _discovery_completo()
        d["costos_finops"] = ""
        res = ensamblar_discovery(d, caso_de_uso="Caso X")
        self.assertIn("Costos preliminares (FinOps)", res["secciones_omitidas"])
        self.assertNotIn("## Costos preliminares (FinOps)", res["markdown"])
        self.assertIn("Información no capturada", res["markdown"])

    def test_diagramas_mermaid(self):
        diagramas = generar_diagramas(_discovery_completo(), "Caso X")
        self.assertIn("```mermaid", diagramas["flujo_datos.md"])
        self.assertIn("SQLPROD01", diagramas["flujo_datos.md"])
        self.assertIn("Pagos", diagramas["datamesh.md"])
        self.assertIn("erDiagram", diagramas["entidades_preliminar.md"])

    def test_tablas_y_trazabilidad(self):
        tablas = generar_tablas(_discovery_completo(),
                                {"decisiones_pendientes": [
                                    {"decision": "Definir corte", "decisor_sugerido": "PO",
                                     "urgencia": "alta"}]})
        self.assertIn("TRX", tablas["inventario_fuentes.md"])
        self.assertIn("Definir corte", tablas["matriz_decisiones.md"])
        self.assertIn("E001", tablas["matriz_trazabilidad.md"])


# =============================================================================================
# v3 — Dos vistas, temas faltantes, fidelidad y export docx
# =============================================================================================
class TestDosVistas(unittest.TestCase):
    def test_documento_tiene_parte_general_y_tecnica(self):
        md = ensamblar_discovery(_discovery_completo(), caso_de_uso="Caso X")["markdown"]
        self.assertIn("Parte I — Discovery inicial (negocio)", md)
        self.assertIn("Parte II — Discovery técnico (datos y habilitación)", md)
        self.assertIn("Cierre — Evaluación", md)

    def test_fuentes_muestran_columnas_y_librerias(self):
        md = ensamblar_discovery(_discovery_completo(), caso_de_uso="Caso X")["markdown"]
        self.assertIn("Columnas", md)       # encabezado de tabla de fuentes
        self.assertIn("Librerias", md)
        self.assertIn("pyodbc", md)


class TestTemasFaltantes(unittest.TestCase):
    def test_schema_vacio_genera_temas_faltantes(self):
        res = ensamblar_discovery(schema_vacio(), caso_de_uso="Caso X")
        self.assertGreater(len(res["temas_faltantes"]), 0)
        self.assertIn("Temas faltantes y deuda técnica", res["markdown"])

    def test_falta_objetivo_aparece_en_temas(self):
        d = _discovery_completo()
        d["objetivo"] = ""   # se desconoce el objetivo
        res = ensamblar_discovery(d, caso_de_uso="Caso X")
        puntos = " ".join(t["punto"] for t in res["temas_faltantes"]).lower()
        self.assertIn("objetivo", puntos)

    def test_supuesto_sin_validar_es_deuda(self):
        score = score_discovery(_discovery_completo())
        d = _discovery_completo()
        d["supuestos"] = [{"descripcion": "Volumetría asumida", "validado_por": ""}]
        temas = recolectar_temas_faltantes(score, d, None)
        tipos = {t["tipo"] for t in temas}
        self.assertIn("supuesto_no_validado", tipos)

    def test_inconsistencias_se_integran(self):
        score = score_discovery(_discovery_completo())
        inc = {"vacios_no_capturados_por_rubrica": ["Falta política de retención"],
               "decisiones_pendientes": [{"decision": "Definir corte", "decisor_sugerido": "PO",
                                          "urgencia": "alta"}]}
        temas = recolectar_temas_faltantes(score, _discovery_completo(), inc)
        puntos = " ".join(t["punto"] for t in temas)
        self.assertIn("retención", puntos)
        self.assertIn("Definir corte", puntos)


class TestFidelidad(unittest.TestCase):
    def test_campo_sin_cita_se_detecta(self):
        d = _discovery_completo()
        d["necesidad_negocio"] = "Resolver reportería manual"   # sin [E00x]
        fid = verificar_fidelidad(d, redaccion_md=None, index=None)
        etiquetas = " ".join(fid["campos_sin_cita"])
        self.assertIn("Problema de negocio", etiquetas)
        self.assertLess(fid["cobertura_pct"], 100.0)

    def test_cita_invalida_contra_indice(self):
        d = _discovery_completo()
        d["necesidad_negocio"] = "Algo respaldado [E999]"
        index = {"items": [{"id": "E001"}]}
        fid = verificar_fidelidad(d, redaccion_md=None, index=index)
        self.assertTrue(any(c["cita"].startswith("E999") for c in fid["citas_invalidas"]))

    def test_parrafo_narrativa_sin_cita(self):
        narr = ("Este es un párrafo largo que afirma un hecho de negocio sin citar ninguna "
                "evidencia que lo respalde, por lo que debería marcarse como no anclado.")
        fid = verificar_fidelidad(_discovery_completo(), redaccion_md=narr, index=None)
        self.assertEqual(len(fid["parrafos_sin_cita"]), 1)


class TestDocxExport(unittest.TestCase):
    def test_genera_docx_con_partes(self):
        if not docx_disponible():
            self.skipTest("python-docx no instalado")
        from docx import Document
        out = os.path.join(_TMP, "demo.docx")
        d = _discovery_completo()
        fid = verificar_fidelidad(d, redaccion_md=None, index={"items": [{"id": "E001"}]})
        ruta = exportar_docx(out, d, caso="Caso X",
                             redaccion={"secciones": {"contexto": "Narrativa [E001]."}},
                             fidelidad=fid, diagramas=generar_diagramas(d, "Caso X"))
        self.assertEqual(ruta, out)
        documento = Document(out)
        texto = "\n".join(p.text for p in documento.paragraphs)
        self.assertIn("Parte I — Discovery inicial (negocio)", texto)
        self.assertIn("Parte II — Discovery técnico (datos y habilitación)", texto)
        self.assertIn("Cobertura de citas", texto)
        self.assertIn("Narrativa", texto)

    def test_docx_tiene_portada_indice_y_estilos(self):
        if not docx_disponible():
            self.skipTest("python-docx no instalado")
        from docx import Document
        out = os.path.join(_TMP, "demo_portada.docx")
        d = _discovery_completo()
        exportar_docx(out, d, caso="Caso X")
        documento = Document(out)
        texto = "\n".join(p.text for p in documento.paragraphs)
        # Portada: el veredicto se ve antes que cualquier sección.
        self.assertIn("DISCOVERY DE PRODUCTO DE DATOS", texto)
        self.assertIn("Veredicto: Go", texto)
        self.assertIn("Contenido", texto)          # título del índice
        # Estilos propios creados por render/estilos_word.py
        nombres = [s.name for s in documento.styles]
        for estilo in ("DiscoveryCuerpo", "DiscoveryDato", "DiscoveryNota", "DiscoveryCita"):
            self.assertIn(estilo, nombres)
        # El índice y la numeración de páginas son campos de Word, no texto fijo.
        xml = documento.element.xml
        self.assertIn("TOC", xml)
        self.assertIn("NUMPAGES", documento.sections[0].footer.paragraphs[0]._p.xml)

    def test_docx_describe_los_diagramas_en_prosa(self):
        if not docx_disponible():
            self.skipTest("python-docx no instalado")
        from docx import Document
        out = os.path.join(_TMP, "demo_diagramas.docx")
        d = _discovery_completo()
        exportar_docx(out, d, caso="Caso X", diagramas=generar_diagramas(d, "Caso X"))
        texto = "\n".join(p.text for p in Document(out).paragraphs)
        self.assertIn("Anexos", texto)
        self.assertIn("SQLPROD01", texto)
        self.assertIn("medallón", texto)   # lectura en prosa del flujo, no solo el código


# =============================================================================================
# Básicos que se conservan
# =============================================================================================
class TestBasicos(unittest.TestCase):
    def test_manifest_clasifica(self):
        m = build_manifest(DEMO)
        self.assertEqual(m["total_items"], 3)
        self.assertTrue(all(i["modificado"] for i in m["items"]))  # metadatos presentes

    def test_excel_detecta_fuentes(self):
        data = extract_excel(os.path.join(DEMO, "fuentes", "fuentes_volumetria.xlsx"))
        self.assertEqual(len(data["fuentes_detectadas"]), 3)

    def test_todos_los_prompts_cargan(self):
        nombres = listar_prompts()
        self.assertIn("redactor-experto", nombres)
        for n in nombres:
            self.assertTrue(cargar_prompt(n).strip())


# =============================================================================================
# Evolución v3 — campos nuevos, narrativa-first, HTML ejecutivo, ensamblado unificado, formatos
# =============================================================================================
class TestCamposNuevos(unittest.TestCase):
    def test_schema_incluye_campos_obligatorios(self):
        d = schema_vacio()
        for k in ("casos_uso", "fecha_discovery", "contexto", "problematica_actual",
                  "product_owner", "lider_tecnico", "beneficio_economico",
                  "hallazgos_tecnicos", "productos_relacionados", "mensajes_clave"):
            self.assertIn(k, d)

    def test_po_sin_nombre_no_da_go(self):
        d = _discovery_completo()
        d["product_owner"] = {"nombre": "", "rol": "", "contacto": ""}
        self.assertNotEqual(score_discovery(d)["veredicto"], "Go")

    def test_beneficio_economico_sin_evidencia_no_se_afirma(self):
        d = _discovery_completo()
        d["beneficio_economico"] = "Ahorra mucho dinero"  # sin [E00x]
        md = ensamblar_discovery(d, caso_de_uso="X")["markdown"]
        self.assertIn("No cuantificado con evidencia", md)


class TestNarrativaFirst(unittest.TestCase):
    def test_dependencias_como_prosa_no_tabla(self):
        d = _discovery_completo()
        salida = render_clave("dependencias", d["dependencias"])
        self.assertNotIn("|", salida)          # no es tabla markdown
        self.assertIn("**", salida)            # es prosa con enunciado en negrita

    def test_fuentes_si_es_tabla(self):
        d = _discovery_completo()
        salida = render_clave("fuentes", d["fuentes"])
        self.assertIn("|", salida)             # el inventario de fuentes SÍ es tabla

    def test_matrices_auxiliares_fuerzan_tabla(self):
        d = _discovery_completo()
        tablas = generar_tablas(d)
        self.assertIn("|", tablas["matriz_dependencias.md"])  # el archivo de matriz sí tabula


class TestHtmlExport(unittest.TestCase):
    def test_html_standalone_sin_refs_externas(self):
        d = _discovery_completo()
        d["nombre_caso_uso"] = "Caso Demo"
        html = render_html(d, caso="Caso Demo", fidelidad={"cobertura_pct": 100.0})
        self.assertNotIn("{{", html)                       # todos los tokens reemplazados
        for ref in ("http://", "https://", "<link", "cdn"):
            self.assertNotIn(ref, html)
        self.assertIn("<!doctype html>", html.lower())
        self.assertIn("Caso Demo", html)

    def test_html_veredicto_clase_coincide(self):
        d = _discovery_completo()  # completo → Go
        html = render_html(d, caso="X")
        self.assertIn('class="v-go"', html)

    def test_html_escapa_datos(self):
        d = _discovery_completo()
        d["mensajes_clave"] = ["<script>alert(1)</script>"]
        html = render_html(d, caso="X")
        self.assertNotIn("<script>alert(1)", html)         # escapado
        self.assertIn("&lt;script&gt;", html)


class TestEntregablesUnificado(unittest.TestCase):
    CASO = "caso_unif"

    def setUp(self):
        d = _discovery_completo()
        d["fecha_discovery"] = ""   # vacío → debe derivarse del manifiesto
        workspace.save_json(self.CASO, "30_discovery.json", d)
        workspace.save_json(self.CASO, "00_manifest.json",
                            {"items": [{"modificado": "2026-07-20T10:00:00"}]})
        workspace.save_json(self.CASO, os.path.join("10_evidence", "index.json"),
                            {"items": [{"id": "E001"}]})

    def test_genera_md_html_y_fidelidad(self):
        r = construir_entregables(self.CASO, caso_de_uso="Caso Unificado")
        self.assertNotIn("error", r)
        base = os.path.join(_TMP, self.CASO, "outputs")
        self.assertTrue(os.path.exists(os.path.join(base, "discovery_final.md")))
        self.assertTrue(os.path.exists(os.path.join(base, "resumen_ejecutivo.html")))
        self.assertTrue(os.path.exists(os.path.join(_TMP, self.CASO, "65_fidelidad.json")))
        if docx_disponible():
            self.assertTrue(os.path.exists(os.path.join(base, "discovery_final.docx")))

    def test_fecha_se_deriva_del_manifest(self):
        construir_entregables(self.CASO, caso_de_uso="X")
        html = workspace.load_text(self.CASO, "outputs/resumen_ejecutivo.html")
        self.assertIn("2026-07-20", html)   # fecha rellenada desde el manifiesto


class TestFormatos(unittest.TestCase):
    def test_manifest_avisa_de_legacy(self):
        tmp = tempfile.mkdtemp(prefix="fmt_")
        os.makedirs(os.path.join(tmp, "fuentes"))
        open(os.path.join(tmp, "fuentes", "viejo.xls"), "w").close()
        m = build_manifest(tmp)
        self.assertTrue(any("legacy" in a.lower() for a in m["advertencias"]))
        self.assertEqual(m["total_items"], 0)  # el .xls no se cuenta como procesable
        shutil.rmtree(tmp, ignore_errors=True)

    def test_csv_se_extrae(self):
        tmp = tempfile.mkdtemp(prefix="csv_")
        ruta = os.path.join(tmp, "datos.csv")
        with open(ruta, "w", encoding="utf-8") as fh:
            fh.write("tabla,servidor\nTRX,SQLPROD01\n")
        data = extract_document(ruta)
        self.assertIn("TRX", data["texto"])
        self.assertIn("SQLPROD01", data["texto"])
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=2)


# =============================================================================================
# v4 — Formatos nuevos, deduplicación, validación y sincronía de la documentación
# =============================================================================================
class TestRegistroDeFormatos(unittest.TestCase):
    def test_toda_extension_tiene_adaptador(self):
        """El registro es la única tabla: no puede clasificar algo que nadie sabe procesar."""
        from discovery_agents.ingest.registry import FORMATOS, adaptador
        for formato in FORMATOS:
            self.assertTrue(callable(adaptador(formato.extractor)), formato.extension)

    def test_extractor_desconocido_da_error_util(self):
        from discovery_agents.ingest.registry import adaptador
        with self.assertRaises(ValueError) as ctx:
            adaptador("inventado")
        self.assertIn("Categorías registradas", str(ctx.exception))

    def test_informe_de_formatos(self):
        from discovery_agents.ingest.registry import formatos_soportados
        informe = formatos_soportados()
        self.assertGreater(informe["total_extensiones"], 40)
        for categoria in ("excel", "documento", "estructurado", "correo", "imagen", "legacy"):
            self.assertIn(categoria, informe["por_extractor"])


class TestExtractoresNuevos(unittest.TestCase):
    def _escribir(self, nombre, contenido):
        ruta = os.path.join(_TMP, nombre)
        with open(ruta, "w", encoding="utf-8") as fh:
            fh.write(contenido)
        return ruta

    def test_sql_aporta_tablas_al_inventario(self):
        from discovery_agents.ingest.extractores.estructurados import extraer
        ruta = self._escribir("carga.sql", """-- comentario con FROM que no cuenta
            SELECT a.id FROM BACPROD.dbo.TRX_ADQUIRENCIA a
            JOIN dbo.CATALOGO_COMERCIOS b ON a.cid = b.id;""")
        registro, resumen, extras = extraer(ruta)
        tablas = [t["tabla"] for t in extras["fuentes"]]
        self.assertIn("TRX_ADQUIRENCIA", tablas)
        self.assertIn("CATALOGO_COMERCIOS", tablas)
        self.assertIn("SQL", resumen)
        self.assertIn("vista_texto", registro)

    def test_json_se_resume_por_estructura(self):
        from discovery_agents.ingest.extractores.estructurados import extraer
        ruta = self._escribir("conf.json",
                              '{"origen":{"servidor":"SQLPROD01","tablas":[{"nombre":"TRX"}]}}')
        registro, _, _ = extraer(ruta)
        self.assertIn("origen.servidor: SQLPROD01", registro["vista_texto"])
        self.assertIn("origen.tablas[0].nombre: TRX", registro["vista_texto"])

    def test_correo_extrae_cabeceras_y_poda_el_hilo(self):
        from discovery_agents.ingest.extractores.correo import extraer
        ruta = self._escribir("acuerdo.eml", "From: Carlos Vega <cv@x.com>\n"
                              "Subject: Acuerdo de replica\n"
                              "Date: Mon, 14 Jul 2026 10:00:00 -0600\n\n"
                              "Confirmamos la replica no transaccional.\n\n"
                              "> texto citado del hilo anterior\n")
        registro, resumen, _ = extraer(ruta)
        self.assertIn("Carlos Vega", resumen)
        self.assertIn("Acuerdo de replica", resumen)
        self.assertIn("Confirmamos la replica", registro["vista_texto"])
        self.assertNotIn("texto citado del hilo", registro["vista_texto"])

    def test_html_se_extrae_sin_marcado(self):
        from discovery_agents.ingest.extractores.documentos import extract_document
        ruta = self._escribir("pagina.html",
                              "<html><head><style>.x{color:red}</style></head>"
                              "<body><h1>Fuentes</h1><p>La tabla TRX vive en SQLPROD01</p>"
                              "<script>var a=1</script></body></html>")
        texto = extract_document(ruta)["texto"]
        self.assertIn("La tabla TRX vive en SQLPROD01", texto)
        self.assertNotIn("color:red", texto)
        self.assertNotIn("var a=1", texto)

    def test_svg_se_lee_sin_ocr(self):
        from discovery_agents.ingest.extractores.imagenes import extract_image
        ruta = self._escribir("diagrama.svg",
                              '<svg xmlns="http://www.w3.org/2000/svg">'
                              '<text>SQLPROD01</text><text>Azure Data Factory</text></svg>')
        data = extract_image(ruta)
        self.assertIn("SQLPROD01", data["texto_ocr"])
        self.assertFalse(data["requiere_vision_llm"])   # el XML no necesita modelo de visión

    def test_legacy_sin_libreoffice_avisa_en_vez_de_romper(self):
        from discovery_agents.ingest.extractores.legacy import ConversionNoDisponible, convertir
        from discovery_agents.ingest.registry import soffice_ruta
        if soffice_ruta():
            self.skipTest("LibreOffice instalado: la conversión sí está disponible")
        ruta = self._escribir("viejo.xls", "contenido binario simulado")
        with self.assertRaises(ConversionNoDisponible) as ctx:
            convertir(ruta)
        self.assertIn("LibreOffice", str(ctx.exception))

    def test_imagenes_incrustadas_se_listan(self):
        """Un .docx/.pptx es un ZIP: sus medios se localizan sin librería alguna."""
        import zipfile
        from discovery_agents.ingest.extractores.embebidos import listar_incrustadas
        ruta = os.path.join(_TMP, "presentacion.pptx")
        with zipfile.ZipFile(ruta, "w") as z:
            z.writestr("ppt/media/image1.png", b"\x89PNG" + b"0" * 20000)   # supera el mínimo
            z.writestr("ppt/media/icono.png", b"\x89PNG" + b"0" * 100)      # icono: se ignora
            z.writestr("ppt/slides/slide1.xml", "<xml/>")
        incrustadas = listar_incrustadas(ruta)
        self.assertEqual(incrustadas, ["ppt/media/image1.png"])


class TestDeduplicacion(unittest.TestCase):
    def test_funde_reformulaciones_y_une_procedencia(self):
        from discovery_agents.knowledge.dedup import deduplicar
        resultado = deduplicar([
            {"hallazgo": "La replica evita bloqueos en produccion",
             "evidencia": "E001", "segmentos": ["seg-004"]},
            {"hallazgo": "Las replicas evitan el bloqueo en produccion",
             "evidencia": "E002", "segmentos": ["seg-011"]},
        ])
        self.assertEqual(resultado["fusionados"], 1)
        self.assertEqual(len(resultado["items"]), 1)
        # La procedencia se GANA al fusionar, no se pierde. Cada campo conserva su tipo: la
        # evidencia es un string en el schema y sigue siéndolo; los segmentos son una lista.
        fundido = resultado["items"][0]
        self.assertIn("E001", fundido["evidencia"])
        self.assertIn("E002", fundido["evidencia"])
        self.assertEqual(sorted(fundido["segmentos"]), ["seg-004", "seg-011"])
        self.assertTrue(resultado["grupos"])   # la fusión queda auditable

    def test_no_funde_hechos_distintos_con_vocabulario_comun(self):
        """El error grave sería fundir dos tablas distintas por compartir la frase."""
        from discovery_agents.knowledge.dedup import deduplicar
        resultado = deduplicar([
            {"hallazgo": "La tabla TRX tiene 400M registros"},
            {"hallazgo": "La tabla CATALOGO tiene 400M registros"},
        ])
        self.assertEqual(resultado["fusionados"], 0)
        self.assertEqual(len(resultado["items"]), 2)

    def test_raiz_es_consistente_entre_singular_y_plural(self):
        from discovery_agents.common.texto import raiz
        for singular, plural in [("dato", "datos"), ("tabla", "tablas"), ("riesgo", "riesgos"),
                                 ("servidor", "servidores"), ("fuente", "fuentes"),
                                 ("pesada", "pesadas"), ("dependencia", "dependencias")]:
            self.assertEqual(raiz(singular), raiz(plural), f"{singular}/{plural}")


class TestValidadorDeForma(unittest.TestCase):
    def test_lista_aplanada_a_texto_es_error(self):
        from discovery_agents.knowledge.validador import validar_discovery
        d = _discovery_completo()
        d["fuentes"] = "TRX en SQLPROD01, LIQ en SQLPROD02"
        informe = validar_discovery(d)
        self.assertFalse(informe["valido"])
        self.assertTrue(any("fuentes" in e for e in informe["errores"]))

    def test_persona_aplanada_es_error(self):
        from discovery_agents.knowledge.validador import validar_discovery
        d = _discovery_completo()
        d["dueno_producto"] = "Carlos Vega"
        self.assertFalse(validar_discovery(d)["valido"])

    def test_discovery_correcto_es_valido(self):
        from discovery_agents.knowledge.validador import validar_discovery
        self.assertTrue(validar_discovery(_discovery_completo())["valido"])

    def test_ensamblado_se_detiene_con_discovery_roto(self):
        from discovery_agents.render.entregables import construir_entregables
        d = _discovery_completo()
        d["riesgos"] = "hay riesgos de bloqueo"
        workspace.save_json("caso_roto", "30_discovery.json", d)
        resultado = construir_entregables("caso_roto")
        self.assertIn("error", resultado)
        self.assertTrue(resultado["detalles"])


class TestCoherenciaDelVeredicto(unittest.TestCase):
    def test_desconocido_no_cuenta_como_respondido(self):
        """Un Discovery donde nadie supo nada no puede puntuar como completo."""
        d = _discovery_completo()
        for clave in ("objetivo", "alcance_in", "alcance_out", "metricas_exito"):
            d[clave] = "DESCONOCIDO"
        score = score_discovery(d)
        self.assertNotEqual(score["veredicto"], "Go")
        self.assertGreater(score["n_brechas"], 0)

    def test_critica_del_critico_llega_a_temas_faltantes(self):
        """50_critica.json se escribía y nadie lo leía: su trabajo no llegaba al documento."""
        d = _discovery_completo()
        critica = {"brechas_priorizadas": [{
            "bloque": "Fuentes", "que_falta": "ventana temporal",
            "por_que_importa": "cambia el dimensionamiento",
            "pregunta_poderosa": "¿Cuál es la ventana temporal de la volumetría declarada?",
            "severidad": "alta"}]}
        res = ensamblar_discovery(d, caso_de_uso="Caso X", critica=critica)
        puntos = " ".join(t["punto"] for t in res["temas_faltantes"])
        self.assertIn("ventana temporal", puntos)
        self.assertIn("ventana temporal", res["markdown"])

    def test_riesgos_vacios_se_declaran_sin_inventar(self):
        d = _discovery_completo()
        d["riesgos"] = []
        md = ensamblar_discovery(d, caso_de_uso="Caso X")["markdown"]
        self.assertIn("No se identificaron riesgos", md)
        self.assertIn("no garantía de su inexistencia", md)

    def test_riesgos_se_agrupan_por_tipo(self):
        d = _discovery_completo()
        d["riesgos"] = [
            {"descripcion": "Bloqueos en produccion", "tipo": "tecnico", "severidad": "alta"},
            {"descripcion": "Duplicados en el origen", "tipo": "calidad", "severidad": "media"},
        ]
        md = ensamblar_discovery(d, caso_de_uso="Caso X")["markdown"]
        self.assertIn("Técnicos", md)
        self.assertIn("De calidad de datos", md)


class TestHtmlEjecutivoV4(unittest.TestCase):
    def test_incluye_completitud_por_bloque(self):
        html = render_html(_discovery_completo(), caso="Caso X")
        self.assertIn("Completitud por bloque", html)
        self.assertIn("Umbral de Go", html)
        self.assertIn("barra-fila", html)

    def test_incluye_gobierno_y_beneficio(self):
        html = render_html(_discovery_completo(), caso="Caso X")
        self.assertIn("Gobierno del producto", html)
        self.assertIn("Beneficio esperado", html)
        self.assertIn("Carlos Vega", html)

    def test_tiene_hoja_de_impresion_y_respeta_movimiento_reducido(self):
        html = render_html(_discovery_completo(), caso="Caso X")
        self.assertIn("@media print", html)
        self.assertIn("prefers-reduced-motion", html)

    def test_define_los_dos_temas_explicitamente(self):
        """El interruptor del lector debe poder ganarle a la preferencia del sistema."""
        html = render_html(_discovery_completo(), caso="Caso X")
        self.assertIn('data-theme="dark"', html)
        self.assertIn('data-theme="light"', html)
        self.assertIn("prefers-color-scheme: dark", html)


class TestDocumentacionSincronizada(unittest.TestCase):
    """El Blueprint promete bastar para reconstruir el sistema; estas pruebas fijan esa promesa."""

    _RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

    def _blueprint(self):
        with open(os.path.join(self._RAIZ, "docs", "BLUEPRINT.md"), encoding="utf-8") as fh:
            return fh.read()

    def test_blueprint_refleja_el_codigo(self):
        """Un Blueprint desactualizado reconstruiría una versión inferior del agente."""
        import subprocess
        proceso = subprocess.run(
            [sys.executable, os.path.join(self._RAIZ, "tools", "sync_docs.py"), "--check"],
            capture_output=True, text=True, cwd=self._RAIZ)
        self.assertEqual(proceso.returncode, 0,
                         f"Blueprint desincronizado. {proceso.stderr}")

    def test_blueprint_embebe_todos_los_modulos_integros(self):
        """Sin el código completo, el documento no basta para reconstruir nada."""
        blueprint = self._blueprint()
        paquete = os.path.join(self._RAIZ, "discovery_agents")
        ausentes = []
        for actual, dirs, archivos in os.walk(paquete):
            dirs[:] = [d for d in dirs if d != "__pycache__"]
            for archivo in archivos:
                if not archivo.endswith(".py"):
                    continue
                ruta = os.path.join(actual, archivo)
                with open(ruta, encoding="utf-8") as fh:
                    fuente = fh.read().rstrip()
                if fuente not in blueprint:
                    ausentes.append(os.path.relpath(ruta, self._RAIZ))
        self.assertEqual(ausentes, [], f"Módulos ausentes o truncados: {ausentes}")

    def test_blueprint_embebe_los_entregables_y_contratos(self):
        """La plantilla, la suite y los prompts también son necesarios para reconstruir."""
        blueprint = self._blueprint()
        for rel in ("discovery_agents/render/assets/executive.html", "tests/test_core.py"):
            with open(os.path.join(self._RAIZ, rel), encoding="utf-8") as fh:
                self.assertIn(fh.read().rstrip(), blueprint, rel)
        for nombre in listar_prompts():
            self.assertIn(cargar_prompt(nombre).rstrip(), blueprint, nombre)

    def test_blueprint_no_depende_del_entorno_donde_se_creo(self):
        """Debe poder leerse en cualquier máquina sin referencias a otra.

        Los términos prohibidos se ensamblan por trozos a propósito: este archivo viaja embebido
        dentro del Blueprint, así que escribirlos enteros aquí haría que la prueba se detectara a
        sí misma y fallara siempre.
        """
        import re
        blueprint = self._blueprint()
        rutas = "/" + "Users" + "/|/" + "home" + "/[a-z]"
        proveedores = "|".join(["chat" + "gpt", "open" + "ai", "anthro" + "pic"])
        historial = "|".join(["historial de gener" + "acion", "prompt orig" + "inal"])
        prohibidos = [
            (rutas, "rutas absolutas de un sistema concreto"),
            (r"(?i)\b(" + proveedores + r")\b", "nombres de proveedores de modelos"),
            ("(?i)(" + historial + ")", "historial de generación"),
        ]
        for patron, descripcion in prohibidos:
            encontrados = re.findall(patron, blueprint)
            self.assertEqual(encontrados, [],
                             f"El Blueprint contiene {descripcion}: {encontrados[:3]}")
